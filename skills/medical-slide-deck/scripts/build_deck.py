#!/usr/bin/env python3
"""Render a non-promotional medical slide deck from a JSON spec.

The spec is JSON so the content is reviewable and diffable before it becomes a
binary that nobody can inspect in a pull request. The builder enforces the
mechanical compliance rules — a data slide cannot render without a citation, the
draft marking is stamped on the title slide — and leaves the judgement calls to
deliverable-quality-review and to you.

    S=skills/medical-slide-deck/scripts/build_deck.py

    python3 $S --example > deck.json      # a worked spec to start from
    python3 $S --spec deck.json --out advisory-board.pptx
    python3 $S --spec deck.json --check   # validate without rendering

Requires python-pptx:  pip install python-pptx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

DRAFT = "DRAFT — NOT FOR EXTERNAL USE. REQUIRES QUALIFIED MEDICAL REVIEW."

EXAMPLE = {
    "title": "Bispecific antibodies in relapsed/refractory myeloma",
    "subtitle": "Advisory board — scientific questions for discussion",
    "presenter": "Medical Affairs",
    "date": "2026-09-12",
    "jurisdiction": "US",
    "approval_status": (
        "NORVANTIB is approved in the US for adult patients with "
        "relapsed/refractory multiple myeloma after ≥ 4 prior lines. "
        "Uses discussed beyond this are investigational."
    ),
    "slides": [
        {
            "type": "section",
            "title": "Why we asked you here",
        },
        {
            "type": "bullets",
            "title": "Objectives for today",
            "bullets": [
                "Understand how you are sequencing BCMA-directed therapy in practice",
                "Test whether our proposed evidence plan answers questions you have",
                "Identify what we are missing",
            ],
            "notes": "Say explicitly what will be done with the advice.",
        },
        {
            "type": "data",
            "title": "ORR 63.0% in triple-class-exposed patients",
            "design": "Single-arm, phase 1/2 (n=165). No comparative inference.",
            "rows": [
                ["Endpoint", "Result (95% CI)", "Note"],
                ["Overall response rate", "63.0% (55.2–70.4)", "Primary endpoint"],
                ["Median DOR", "18.4 months (14.9–NE)", "Immature"],
                ["CRS, any grade", "72.1%", "Grade ≥3: 0.6%"],
            ],
            "citation": "Example A, et al. J Example Med. 2024;12(3):100-110. PMID 00000000",
            "tier": "peer-reviewed",
        },
        {
            "type": "bullets",
            "title": "What remains unknown",
            "bullets": [
                "No randomised comparison against current standard of care",
                "No data beyond 24 months of follow-up",
                "Sequencing after BCMA-directed therapy: no prospective data",
            ],
            "notes": "This slide builds credibility. Do not cut it for time.",
        },
        {
            "type": "questions",
            "title": "Questions for the group",
            "questions": [
                "What would you need to see before using this earlier in the pathway?",
                "Where does your practice diverge from the label, and why?",
                "What question would you want answered that we have not asked?",
            ],
        },
        {"type": "references", "title": "References", "references": []},
    ],
}


def validate(spec: dict) -> list[str]:
    """Mechanical checks only. Promotional framing is a human judgement."""
    problems: list[str] = []
    if not spec.get("title"):
        problems.append("spec has no title")
    slides = spec.get("slides") or []
    if not slides:
        problems.append("spec has no slides")

    for i, s in enumerate(slides, 1):
        kind = s.get("type", "")
        where = f"slide {i} ({kind or 'no type'})"
        if kind not in {"section", "bullets", "data", "questions", "references", "image"}:
            problems.append(f"{where}: unknown slide type")
        if not s.get("title"):
            problems.append(f"{where}: no title")

        if kind == "data":
            # The rule this script exists to enforce.
            if not s.get("citation"):
                problems.append(
                    f"{where}: data slide has no citation. Every data slide carries "
                    f"its source — a slide that cannot be traced cannot be reviewed."
                )
            if not s.get("design"):
                problems.append(
                    f"{where}: data slide does not name the study design. 'Single-arm' "
                    f"and 'randomised' must be visible where the number is, not only "
                    f"in the backup."
                )
            if not s.get("rows"):
                problems.append(f"{where}: data slide has no rows")

    if not spec.get("approval_status"):
        problems.append(
            "spec has no approval_status. State what is approved, in which "
            "jurisdiction, and that anything else is investigational."
        )
    return problems


def build(spec: dict, out: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        raise SystemExit(
            "python-pptx is required.\n  pip install python-pptx"
        )

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)

    BLANK = prs.slide_layouts[6]
    INK = RGBColor(0x1A, 0x1A, 0x1A)
    MUTED = RGBColor(0x5A, 0x5A, 0x5A)
    WARN = RGBColor(0x99, 0x33, 0x00)

    def textbox(slide, left, top, width, height, text, size, *,
                bold=False, colour=INK, align=PP_ALIGN.LEFT, wrap=True):
        box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                       Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = wrap
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = colour
        return tf

    def footer(slide, citation: str = "", tier: str = ""):
        parts = [p for p in [citation, f"[{tier}]" if tier and tier != "peer-reviewed" else ""] if p]
        if parts:
            textbox(slide, 0.6, 6.75, 12.1, 0.5, "  ".join(parts), 9, colour=MUTED)

    def draft_stripe(slide):
        textbox(slide, 0.6, 7.05, 12.1, 0.35, DRAFT, 8, colour=WARN)

    # ---- title slide ----------------------------------------------------
    s = prs.slides.add_slide(BLANK)
    textbox(s, 0.8, 2.2, 11.7, 1.4, spec.get("title", ""), 40, bold=True)
    if spec.get("subtitle"):
        textbox(s, 0.8, 3.6, 11.7, 0.8, spec["subtitle"], 20, colour=MUTED)
    meta = " · ".join(
        x for x in [spec.get("presenter", ""), spec.get("date", ""),
                    spec.get("jurisdiction", "")] if x
    )
    textbox(s, 0.8, 4.5, 11.7, 0.5, meta, 13, colour=MUTED)
    textbox(s, 0.8, 5.3, 11.7, 1.0, spec.get("approval_status", ""), 11, colour=INK)
    textbox(s, 0.8, 6.6, 11.7, 0.5, DRAFT, 12, bold=True, colour=WARN)

    # ---- content slides -------------------------------------------------
    for spec_slide in spec.get("slides", []):
        kind = spec_slide["type"]
        s = prs.slides.add_slide(BLANK)

        if kind == "section":
            textbox(s, 0.8, 3.0, 11.7, 1.2, spec_slide["title"], 32, bold=True)
            draft_stripe(s)
            continue

        textbox(s, 0.6, 0.45, 12.1, 1.0, spec_slide["title"], 26, bold=True)
        top = 1.5

        if spec_slide.get("design"):
            textbox(s, 0.6, 1.45, 12.1, 0.45, spec_slide["design"], 13, colour=MUTED)
            top = 2.1

        if kind == "bullets":
            tf = textbox(s, 0.9, top, 11.5, 4.4, "", 18)
            tf.clear()
            for j, b in enumerate(spec_slide.get("bullets", [])):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = "•  " + b
                p.font.size = Pt(18)
                p.space_after = Pt(14)

        elif kind == "questions":
            tf = textbox(s, 0.9, top, 11.5, 4.4, "", 20)
            tf.clear()
            for j, q in enumerate(spec_slide.get("questions", []), 1):
                p = tf.paragraphs[0] if j == 1 else tf.add_paragraph()
                p.text = f"{j}.  {q}"
                p.font.size = Pt(20)
                p.space_after = Pt(20)

        elif kind == "data":
            rows = spec_slide["rows"]
            n_rows, n_cols = len(rows), len(rows[0])
            table = s.shapes.add_table(
                n_rows, n_cols, Inches(0.7), Inches(top),
                Inches(11.9), Inches(0.45 * n_rows)
            ).table
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = str(val)
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(14 if r else 14)
                            run.font.bold = r == 0

        elif kind == "references":
            tf = textbox(s, 0.9, top, 11.5, 4.6, "", 11)
            tf.clear()
            refs = spec_slide.get("references", [])
            if not refs:
                tf.paragraphs[0].text = (
                    "No references listed. Every data slide must cite its source; "
                    "collect them here."
                )
                tf.paragraphs[0].font.size = Pt(12)
            for j, ref in enumerate(refs):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = f"{j + 1}. {ref}"
                p.font.size = Pt(11)
                p.space_after = Pt(6)

        elif kind == "image":
            path = Path(spec_slide["path"])
            if path.exists():
                s.shapes.add_picture(str(path), Inches(1.4), Inches(top), height=Inches(4.4))
            else:
                textbox(s, 0.9, top, 11.5, 0.6, f"[missing image: {path}]", 14, colour=WARN)

        footer(s, spec_slide.get("citation", ""), spec_slide.get("tier", ""))
        draft_stripe(s)

        if spec_slide.get("notes"):
            s.notes_slide.notes_text_frame.text = spec_slide["notes"]

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--spec", help="JSON deck specification")
    ap.add_argument("--out", default="deck.pptx")
    ap.add_argument("--check", action="store_true", help="validate only, do not render")
    ap.add_argument("--example", action="store_true", help="print a worked spec")
    args = ap.parse_args()

    if args.example:
        print(json.dumps(EXAMPLE, indent=2))
        return 0

    if not args.spec:
        ap.print_help()
        return 2

    spec_path = Path(args.spec)
    if not spec_path.exists():
        print(f"No such file: {spec_path}", file=sys.stderr)
        return 2

    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        print(f"Spec is not valid JSON: {exc}", file=sys.stderr)
        return 2

    problems = validate(spec)
    for p in problems:
        print(f"  BLOCKED  {p}")
    if problems:
        print(
            f"\n{len(problems)} problem(s). These are mechanical compliance rules, "
            f"not style preferences — fix them before rendering.\n"
            f"Note that passing these checks does NOT mean the deck is "
            f"non-promotional. Run deliverable-quality-review for that."
        )
        return 1

    if args.check:
        print("Spec passes the mechanical checks.")
        print(
            "Still required: deliverable-quality-review for promotional framing, "
            "and mlr-review-readiness before the deck goes to review."
        )
        return 0

    out = Path(args.out)
    build(spec, out)
    n = len(spec.get("slides", [])) + 1
    print(f"Wrote {out} ({n} slides).")
    print(
        "\nThe draft marking is stamped on every slide. It is the control that "
        "keeps unreviewed content from reaching an external audience — the "
        "reviewer removes it, not you."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
