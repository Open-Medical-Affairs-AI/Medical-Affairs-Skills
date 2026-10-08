#!/usr/bin/env python3
"""ma_render — the library's standalone deliverable engine.

One file, copied verbatim into every skill's ``scripts/`` folder, so a skill
loaded on its own (a GrokBot agent, a Claude project, a zip upload) can still
produce a finished Word document, PowerPoint deck or PDF. The canonical copy
lives at ``shared/ma_render/ma_render.py``; ``scripts/sync_renderer.py`` copies
it into each skill and CI fails if a copy drifts.

Deliverables are .pptx, .docx and .pdf. Markdown is a drafting format and a
review aid, never the thing handed to a stakeholder.

    R=scripts/ma_render.py                        # inside any skill folder

    python3 $R bootstrap                          # install python deps, report tools
    python3 $R example deck   > deck.json         # worked specs to start from
    python3 $R example report > report.json

    python3 $R deck   deck.json   --out out/deck.pptx     # + deck.pdf when possible
    python3 $R report report.json --out out/report.docx   # + report.pdf
    python3 $R report draft.md    --out out/report.docx   # markdown draft -> Word/PDF
    python3 $R deck   slides.md   --out out/deck.pptx     # '## ' = one slide
    python3 $R qa     out/deck.pptx --preview out/preview # overflow + PNG contact sheet

Missing python packages are installed on first use (python-pptx, python-docx,
matplotlib, reportlab) unless MA_RENDER_NO_INSTALL=1 or --no-install is set.
If installation is impossible the engine still writes a real .docx (standard
library writer) or a print-ready HTML file — never a Markdown deliverable — and
says exactly what was degraded and how to get the full version.
"""

from __future__ import annotations

import argparse
import html
import importlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path

VERSION = "1.0.0"
DRAFT = "DRAFT — NOT FOR EXTERNAL USE. REQUIRES QUALIFIED MEDICAL REVIEW."
SYNTHETIC = "SYNTHETIC WORKSHOP DATA — NOT REAL PATIENTS, PRODUCTS OR STUDIES"

# ---------------------------------------------------------------------------
# Visual language (consulting-grade-design). One ink, one accent, oxblood is
# reserved for safety. Every renderer below reads these values.
# ---------------------------------------------------------------------------
INK = "12283A"
SLATE = "5B6B79"
TEAL = "0E7C7B"
GOLD = "C89B3C"
OXBLOOD = "8C2F39"
CLOUD = "E8ECEF"
MIST = "F4F6F8"
TEAL_TINT = "E3F1F0"
GOLD_TINT = "FBF3E4"
OX_TINT = "F7E9EA"
MUTED_BAR = "A9B6C1"
WHITE = "FFFFFF"
ARMS = [INK, TEAL, GOLD, SLATE, "7A9CB8", OXBLOOD]

HEAD_FONT = "Georgia"
BODY_FONT = "Calibri"

CALLOUT_STYLES = {
    "finding": (TEAL_TINT, TEAL, "KEY FINDING"),
    "safety": (OX_TINT, OXBLOOD, "SAFETY"),
    "limitation": (GOLD_TINT, GOLD, "LIMITATION"),
    "note": (MIST, SLATE, "NOTE"),
    "decision": (TEAL_TINT, INK, "DECISION NEEDED"),
}

PACKAGES = {
    "pptx": "python-pptx>=1.0",
    "docx": "python-docx>=1.1",
    "matplotlib": "matplotlib>=3.7",
    "reportlab": "reportlab>=4.0",
}


# ===========================================================================
# Capabilities and dependency bootstrap
# ===========================================================================

def have(module: str) -> bool:
    try:
        importlib.import_module(module)
        return True
    except Exception:
        return False


def _install_allowed(flag: bool | None = None) -> bool:
    if flag is not None:
        return flag
    return os.environ.get("MA_RENDER_NO_INSTALL", "") not in {"1", "true", "yes"}


def ensure(modules: list[str], install: bool | None = None, quiet: bool = False) -> dict[str, bool]:
    """Import-check modules; pip-install the missing ones once if allowed."""
    status = {m: have(m) for m in modules}
    missing = [m for m, ok in status.items() if not ok]
    if missing and _install_allowed(install):
        specs = [PACKAGES.get(m, m) for m in missing]
        if not quiet:
            print(f"  installing: {' '.join(specs)}", file=sys.stderr)
        base = [sys.executable, "-m", "pip", "install", "--quiet",
                "--disable-pip-version-check"]
        for extra in ([], ["--user"], ["--break-system-packages"]):
            try:
                proc = subprocess.run(base + extra + specs, capture_output=True,
                                      text=True, timeout=600)
            except Exception:
                break
            if proc.returncode == 0:
                break
        importlib.invalidate_caches()
        if "--user" and hasattr(sys, "path"):
            try:
                import site
                user_site = site.getusersitepackages()
                if user_site and user_site not in sys.path and Path(user_site).exists():
                    sys.path.append(user_site)
            except Exception:
                pass
        status = {m: have(m) for m in modules}
    return status


def office_binary() -> str | None:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    for path in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
                 r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if Path(path).exists():
            return path
    return None


def capabilities(install: bool | None = None) -> dict:
    status = ensure(list(PACKAGES), install=install, quiet=True)
    return {
        "python": sys.version.split()[0],
        "packages": status,
        "office_converter": office_binary(),
        "pdf_to_png": have("pymupdf") or have("fitz") or bool(shutil.which("pdftoppm")),
        "deliverables": {
            "pptx": status["pptx"],
            "docx": True,  # standard-library writer is the floor
            "pdf": status["reportlab"] or bool(office_binary()),
            "charts": status["matplotlib"],
        },
    }


# ===========================================================================
# Text helpers
# ===========================================================================

_INLINE = re.compile(r"(\*\*[^*]+?\*\*|__[^_]+?__|(?<![\w*])\*(?!\s)[^*]+?(?<!\s)\*(?![\w*])|`[^`]+`)")


def runs(text: str) -> list[tuple[str, bool, bool]]:
    """Split light inline markup (**bold**, *italic*) into (text, bold, italic)."""
    out: list[tuple[str, bool, bool]] = []
    for part in _INLINE.split(str(text)):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (
                part.startswith("__") and part.endswith("__")):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True))
        elif part.startswith("`") and part.endswith("`"):
            out.append((part[1:-1], False, False))
        else:
            out.append((part, False, False))
    return out


def plain(text) -> str:
    return "".join(t for t, _, _ in runs(str(text)))


def wrapped_lines(text: str, chars_per_line: int) -> int:
    lines, cur = 0, 0
    for para in str(text).split("\n"):
        lines += 1
        cur = 0
        for word in para.split():
            n = len(word)
            if cur == 0:
                cur = n
            elif cur + 1 + n <= chars_per_line:
                cur += 1 + n
            else:
                lines += 1
                cur = n
            while cur > chars_per_line:
                lines += 1
                cur -= chars_per_line
    return max(lines, 1)


def text_height(paras: list[str], width_in: float, pt: float, *, em: float = 0.5,
                spacing: float = 1.2, gap_pt: float = 0.0) -> float:
    cpl = max(6, int(width_in * 72 / (pt * em)))
    lines = sum(wrapped_lines(plain(p), cpl) for p in paras)
    return (lines * pt * spacing + gap_pt * max(0, len(paras) - 1)) / 72.0


def fit_size(paras: list[str], width_in: float, height_in: float, max_pt: int,
             min_pt: int, *, em: float = 0.5, spacing: float = 1.2,
             gap_ratio: float = 0.5) -> tuple[int, bool]:
    """Largest font size at which the paragraphs fit the box; (size, overflow)."""
    usable_w = max(0.5, width_in - 0.2)
    usable_h = max(0.2, height_in - 0.1)
    for pt in range(int(max_pt), int(min_pt) - 1, -1):
        if text_height(paras, usable_w, pt, em=em, spacing=spacing,
                       gap_pt=pt * gap_ratio) <= usable_h:
            return pt, False
    return int(min_pt), True


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-")[:60] or "item"


def _plain_shape(shp) -> None:
    """Drop the theme style reference so no default shadow or outline renders."""
    from pptx.oxml.ns import qn
    el = shp._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)


def _alpha(shp, percent: int) -> None:
    """Set solid-fill opacity (0-100) on an autoshape."""
    from pptx.oxml.ns import qn
    from lxml import etree
    clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(percent * 1000)))


def _rgb(hexstr: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr)


# ===========================================================================
# Markdown drafts -> specs. Agents draft naturally in markdown; this turns the
# draft into a designed deliverable instead of shipping the draft itself.
# ===========================================================================

_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_IMG = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
_SOURCE = re.compile(r"^\s*(?:<sub>)?\s*\*{0,2}(?:Source|Sources|Reference)s?\*{0,2}\s*[:：]\s*(.+?)(?:</sub>)?\s*$", re.I)


def _split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def md_blocks(lines: list[str]) -> list[dict]:
    """Parse markdown body lines into report blocks."""
    blocks: list[dict] = []
    para: list[str] = []
    i = 0

    def flush():
        if para:
            text = " ".join(x.strip() for x in para).strip()
            if text:
                blocks.append({"type": "paragraph", "text": text})
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip("\n")
        stripped = line.strip()
        if not stripped or stripped in {"---", "***", "___"}:
            flush(); i += 1; continue
        m = re.match(r"^(#{3,6})\s+(.*)$", stripped)
        if m:
            flush()
            blocks.append({"type": "heading", "level": min(3, len(m.group(1)) - 1),
                           "text": m.group(2).strip()})
            i += 1; continue
        if stripped.startswith("|") and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            flush()
            rows = [_split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            block = {"type": "table", "rows": rows}
            if i < len(lines) and _SOURCE.match(lines[i]):
                block["source"] = _SOURCE.match(lines[i]).group(1).strip()
                i += 1
            blocks.append(block); continue
        if re.match(r"^\s*([-*+]|\d+[.)])\s+", line):
            flush()
            numbered = bool(re.match(r"^\s*\d+[.)]\s+", line))
            items: list = []
            while i < len(lines) and re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i]):
                raw = lines[i]
                indent = len(raw) - len(raw.lstrip())
                text = re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", raw).strip()
                text = re.sub(r"^\[[ xX]\]\s*", "", text)
                if indent >= 2 and items:
                    prev = items[-1]
                    if isinstance(prev, str):
                        items[-1] = {"text": prev, "sub": [text]}
                    else:
                        prev.setdefault("sub", []).append(text)
                else:
                    items.append(text)
                i += 1
                # lazy continuation lines
                while (i < len(lines) and lines[i].strip() and not re.match(
                        r"^\s*([-*+]|\d+[.)]|#|\||>)", lines[i])):
                    cont = lines[i].strip()
                    if isinstance(items[-1], dict):
                        items[-1]["text"] += " " + cont
                    else:
                        items[-1] += " " + cont
                    i += 1
            blocks.append({"type": "numbered" if numbered else "bullets", "items": items})
            continue
        if stripped.startswith(">"):
            flush()
            quote: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            text = " ".join(q for q in quote if q)
            style = "note"
            low = plain(text).lower()
            if any(k in low[:40] for k in ("safety", "adverse", "⚠", "warning", "pqc")):
                style = "safety"
            elif any(k in low[:40] for k in ("limitation", "caveat", "uncertain", "gap")):
                style = "limitation"
            elif any(k in low[:40] for k in ("key finding", "finding", "bottom line", "so what", "takeaway", "insight")):
                style = "finding"
            elif any(k in low[:40] for k in ("decision", "ask", "recommend")):
                style = "decision"
            title = ""
            mt = re.match(r"^\*\*([^*]{1,60}?)[.:]?\*\*[.:]?\s*(.*)$", text)
            if mt:
                title, text = mt.group(1).strip(), mt.group(2).strip()
            blocks.append({"type": "callout", "style": style, "title": title, "text": text})
            continue
        mi = _IMG.match(stripped)
        if mi:
            flush()
            blocks.append({"type": "image", "path": mi.group(2), "caption": mi.group(1)})
            i += 1; continue
        ms = _SOURCE.match(stripped)
        if ms and not para:
            blocks.append({"type": "source", "text": ms.group(1).strip()})
            i += 1; continue
        para.append(stripped)
        i += 1
    flush()
    return blocks


def md_to_report(text: str) -> dict:
    lines = text.splitlines()
    spec: dict = {"title": "", "sections": []}
    body_start = 0
    for idx, line in enumerate(lines):
        m = re.match(r"^#\s+(.*)$", line.strip())
        if m:
            spec["title"] = m.group(1).strip()
            body_start = idx + 1
            break
    # optional italic subtitle directly under the title
    j = body_start
    while j < len(lines) and not lines[j].strip():
        j += 1
    if j < len(lines) and re.match(r"^\*[^*].*\*$|^_[^_].*_$", lines[j].strip()):
        spec["subtitle"] = lines[j].strip().strip("*_ ")
        body_start = j + 1
    current = {"heading": "", "lines": []}
    chunks = []
    for line in lines[body_start:]:
        m = re.match(r"^##\s+(.*)$", line.strip())
        if m:
            chunks.append(current)
            current = {"heading": m.group(1).strip(), "lines": []}
        else:
            current["lines"].append(line)
    chunks.append(current)
    refs: list[str] = []
    for ch in chunks:
        blocks = md_blocks(ch["lines"])
        if not ch["heading"] and not blocks:
            continue
        if re.match(r"^(references|sources|bibliography)$", ch["heading"].strip().lower()):
            for b in blocks:
                if b["type"] in ("numbered", "bullets"):
                    refs += [x if isinstance(x, str) else x["text"] for x in b["items"]]
                elif b["type"] == "paragraph":
                    refs.append(b["text"])
            continue
        if not ch["heading"]:
            # text before the first section: drop DRAFT banners (the renderer adds them)
            blocks = [b for b in blocks if "DRAFT" not in plain(b.get("text", ""))]
            if not blocks:
                continue
        spec["sections"].append({"heading": ch["heading"], "blocks": blocks})
    if refs:
        spec["references"] = refs
    if not spec["title"]:
        spec["title"] = spec["sections"][0]["heading"] if spec["sections"] else "Untitled"
    return spec


def md_to_deck(text: str) -> dict:
    """'# ' is the deck title; each '## ' starts a slide."""
    rep = md_to_report(text)
    deck: dict = {"title": rep.get("title", ""), "subtitle": rep.get("subtitle", ""),
                  "slides": []}
    for sec in rep["sections"]:
        if not sec["heading"]:
            continue
        slide: dict = {"title": sec["heading"]}
        bullets: list = []
        for b in sec["blocks"]:
            t = b["type"]
            if t in ("bullets", "numbered"):
                bullets += b["items"]
            elif t == "paragraph":
                if re.match(r"^(design|kicker)\s*:", plain(b["text"]), re.I):
                    slide["design"] = plain(b["text"]).split(":", 1)[1].strip()
                else:
                    bullets.append(b["text"])
            elif t == "table":
                slide.update({"type": "table", "rows": b["rows"]})
                if b.get("source"):
                    slide["citation"] = b["source"]
            elif t == "callout":
                slide["takeaway"] = (b.get("title") + ": " if b.get("title") else "") + b["text"]
            elif t == "source":
                slide["citation"] = b["text"]
            elif t == "image":
                slide.update({"type": "image", "path": b["path"]})
        if "type" not in slide:
            slide["type"] = "bullets"
        if bullets:
            slide["bullets"] = bullets
        deck["slides"].append(slide)
    if rep.get("references"):
        deck["slides"].append({"type": "references", "title": "References",
                               "references": rep["references"]})
    return deck


def md_to_docx(md_text: str, out: Path) -> Path:
    """Markdown draft -> designed .docx (python-docx) or standard-library .docx."""
    spec = md_to_report(md_text)
    out = Path(out)
    if have("docx"):
        workdir = out.parent / f".{out.stem}_assets"
        workdir.mkdir(parents=True, exist_ok=True)
        ReportBuilder(spec, workdir).build(out)
    else:
        report_docx_stdlib(spec, out)
    return out


def md_to_html(md_text: str) -> str:
    """Markdown draft -> self-contained, print-ready HTML document."""
    return report_html(md_to_report(md_text))


def load_spec(path: Path, kind: str) -> dict:
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() in {".md", ".markdown", ".txt"}:
        return md_to_deck(text) if kind == "deck" else md_to_report(text)
    return json.loads(text)


# ===========================================================================
# Charts (matplotlib). Used for Word/PDF figures and for deck charts that need
# confidence intervals, forest plots or KM curves. Self-contained styling.
# ===========================================================================

NATIVE_KINDS = {"column", "bar", "line", "stacked", "stacked100"}
KIND_ALIASES = {"vbar": "column", "grouped": "column", "hbar": "bar",
                "horizontal": "bar", "trend": "line", "stack": "stacked"}


def norm_chart(chart: dict) -> dict:
    c = dict(chart)
    kind = str(c.get("kind", c.get("type", "column"))).lower()
    c["kind"] = KIND_ALIASES.get(kind, kind)
    if "series" not in c and "values" in c:
        c["series"] = [{"name": c.get("name", c.get("y_label", "Value")),
                        "values": c["values"],
                        "ci_low": c.get("ci_low"), "ci_high": c.get("ci_high")}]
    c.setdefault("series", [])
    c.setdefault("categories", [])
    c.setdefault("decimals", 1)
    hl = c.get("highlight")
    if isinstance(hl, str) and hl in c["categories"]:
        c["highlight"] = c["categories"].index(hl)
    return c


def number_format(c: dict) -> str:
    d = int(c.get("decimals", 1))
    base = "0" + ("." + "0" * d if d > 0 else "")
    unit = c.get("unit", "")
    if unit == "%":
        return base + '"%"'
    if unit:
        return base + ' "' + unit + '"'
    return base


def fmt_value(v: float, c: dict) -> str:
    d = int(c.get("decimals", 1))
    s = f"{v:.{d}f}".replace("-", "−")
    unit = c.get("unit", "")
    return s + ("%" if unit == "%" else (" " + unit if unit else ""))


def _mpl_setup():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Calibri", "Carlito", "Arial", "Liberation Sans", "DejaVu Sans"],
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.edgecolor": "#" + SLATE, "axes.labelcolor": "#" + SLATE,
        "xtick.color": "#" + SLATE, "ytick.color": "#" + SLATE,
        "text.color": "#" + INK, "axes.titlecolor": "#" + INK,
        "axes.grid": True, "grid.color": "#" + CLOUD, "grid.linewidth": 0.8,
        "axes.axisbelow": True, "font.size": 11, "axes.titlesize": 13,
        "figure.facecolor": "white", "savefig.facecolor": "white",
    })
    return plt


def chart_png(chart: dict, out: Path, width_in: float = 10.0, height_in: float = 4.8,
              dpi: int = 200) -> Path:
    plt = _mpl_setup()
    c = norm_chart(chart)
    kind = c["kind"]
    fig, ax = plt.subplots(figsize=(width_in, height_in))
    cats = [str(x) for x in c["categories"]]
    series = c["series"]
    hl = c.get("highlight")

    if kind in ("column", "bar", "stacked", "stacked100"):
        n = max(1, len(series))
        horizontal = kind == "bar"
        idx = list(range(len(cats)))
        width = 0.7 / n if kind in ("column", "bar") else 0.6
        bottoms = [0.0] * len(cats)
        totals = [sum(float(s["values"][k]) for s in series) or 1 for k in idx] if kind == "stacked100" else None
        for si, s in enumerate(series):
            vals = [float(v) for v in s["values"]]
            if totals:
                vals = [v / t * 100 for v, t in zip(vals, totals)]
            if n == 1 and hl is not None:
                colors = ["#" + (TEAL if k == hl else MUTED_BAR) for k in idx]
            elif n == 1:
                colors = ["#" + INK] * len(idx)
            else:
                colors = ["#" + ARMS[si % len(ARMS)]] * len(idx)
            if kind in ("column", "bar"):
                pos = [k - 0.35 + width * (si + 0.5) for k in idx]
                err = None
                if s.get("ci_low") and s.get("ci_high"):
                    err = [[v - float(lo) for v, lo in zip(vals, s["ci_low"])],
                           [float(hi) - v for v, hi in zip(vals, s["ci_high"])]]
                if horizontal:
                    bars = ax.barh(pos, vals, height=width * 0.92, color=colors,
                                   xerr=err, error_kw={"ecolor": "#" + SLATE, "capsize": 3, "lw": 1},
                                   label=s.get("name"))
                else:
                    bars = ax.bar(pos, vals, width=width * 0.92, color=colors,
                                  yerr=err, error_kw={"ecolor": "#" + SLATE, "capsize": 3, "lw": 1},
                                  label=s.get("name"))
                if c.get("labels", True):
                    for b, v in zip(bars, vals):
                        if horizontal:
                            x = b.get_width()
                            ax.annotate(fmt_value(v, c), (x, b.get_y() + b.get_height() / 2),
                                        xytext=(4 if v >= 0 else -4, 0), textcoords="offset points",
                                        ha="left" if v >= 0 else "right", va="center",
                                        fontsize=10, color="#" + INK, fontweight="bold")
                        else:
                            y = b.get_height()
                            ax.annotate(fmt_value(v, c), (b.get_x() + b.get_width() / 2, y),
                                        xytext=(0, 4 if v >= 0 else -12), textcoords="offset points",
                                        ha="center", fontsize=10, color="#" + INK, fontweight="bold")
            else:
                if horizontal:
                    ax.barh(idx, vals, left=bottoms, color=colors, label=s.get("name"))
                else:
                    ax.bar(idx, vals, width=width, bottom=bottoms, color=colors, label=s.get("name"))
                bottoms = [b + v for b, v in zip(bottoms, vals)]
        if horizontal:
            ax.set_yticks(idx, cats)
            ax.invert_yaxis()
            ax.grid(axis="y", visible=False)
            ax.axvline(0, color="#" + SLATE, lw=0.8)
            if c.get("y_label"):
                ax.set_xlabel(c["y_label"])
        else:
            ax.set_xticks(idx, cats)
            ax.grid(axis="x", visible=False)
            ax.axhline(0, color="#" + SLATE, lw=0.8)
            if c.get("y_label"):
                ax.set_ylabel(c["y_label"])
        ax.spines["left"].set_visible(horizontal)
    elif kind == "line":
        for si, s in enumerate(series):
            vals = [float(v) for v in s["values"]]
            col = "#" + ARMS[si % len(ARMS)]
            ax.plot(cats, vals, color=col, lw=2.4, marker="o", ms=5, label=s.get("name"))
            if s.get("ci_low") and s.get("ci_high"):
                ax.fill_between(cats, [float(x) for x in s["ci_low"]],
                                [float(x) for x in s["ci_high"]], color=col, alpha=0.12, lw=0)
            if c.get("labels", True) and len(series) <= 3:
                ax.annotate(f"{s.get('name', '')} {fmt_value(vals[-1], c)}".strip(),
                            (len(cats) - 1, vals[-1]), xytext=(6, 0), textcoords="offset points",
                            va="center", fontsize=10, color=col, fontweight="bold")
        ax.grid(axis="x", visible=False)
        if c.get("y_label"):
            ax.set_ylabel(c["y_label"])
    elif kind == "forest":
        rows = c.get("rows", [])
        ys = list(range(len(rows)))[::-1]
        null = float(c.get("null", 1.0))
        for y, r in zip(ys, rows):
            est, lo, hi = float(r["est"]), float(r["lo"]), float(r["hi"])
            overall = r.get("overall")
            col = "#" + (TEAL if overall else INK)
            ax.plot([lo, hi], [y, y], color=col, lw=1.6)
            ax.plot(est, y, marker="D" if overall else "s", color=col,
                    ms=9 if overall else 6)
            ax.annotate(f"{est:.2f} ({lo:.2f}–{hi:.2f})", (1.02, y),
                        xycoords=("axes fraction", "data"), va="center", fontsize=10,
                        color="#" + INK)
        ax.axvline(null, color="#" + SLATE, lw=1, ls="--")
        ax.set_yticks(ys, [str(r.get("label", "")) + (f"  (n={r['n']})" if r.get("n") else "")
                           for r in rows])
        if c.get("log", True):
            ax.set_xscale("log")
            from matplotlib.ticker import ScalarFormatter
            ax.xaxis.set_major_formatter(ScalarFormatter())
        ax.grid(axis="y", visible=False)
        ax.spines["left"].set_visible(False)
        ax.set_xlabel(c.get("x_label", "Estimate (95% CI)"))
        fav = c.get("favours")
        if fav and len(fav) == 2:
            ax.annotate("← " + fav[0], (0.0, -0.16), xycoords="axes fraction", fontsize=9,
                        color="#" + SLATE)
            ax.annotate(fav[1] + " →", (0.62, -0.16), xycoords="axes fraction", fontsize=9,
                        color="#" + SLATE, ha="right")
        fig.subplots_adjust(right=0.78)
    elif kind == "km":
        for si, arm in enumerate(c.get("arms", [])):
            col = "#" + ARMS[si % len(ARMS)]
            ax.step(arm["times"], arm["surv"], where="post", color=col, lw=2.2,
                    label=arm.get("name"))
        ax.set_ylim(0, 1.02 if max((max(a["surv"]) for a in c.get("arms", [])), default=1) <= 1.0 else None)
        ax.set_xlabel(c.get("x_label", "Months"))
        ax.set_ylabel(c.get("y_label", "Probability"))
        risk = c.get("at_risk")
        if risk:
            txt = "Number at risk\n" + "\n".join(
                f"{name}: " + "   ".join(str(v) for v in vals)
                for name, vals in risk.get("arms", {}).items())
            txt += "\nTime: " + "   ".join(str(t) for t in risk.get("times", []))
            fig.text(0.01, 0.01, txt, fontsize=8.5, color="#" + SLATE, va="bottom")
            fig.subplots_adjust(bottom=0.32)
    else:
        raise ValueError(f"unknown chart kind {kind!r}")

    if int(c.get("decimals", 1)) == 0 and kind in ("column", "bar", "stacked", "line"):
        from matplotlib.ticker import MaxNLocator
        (ax.xaxis if kind == "bar" else ax.yaxis).set_major_locator(MaxNLocator(integer=True))
    if c.get("ref_line") is not None:
        ref = float(c["ref_line"])
        (ax.axvline if kind == "bar" else ax.axhline)(ref, color="#" + GOLD, lw=1.2, ls="--")
    if c.get("y_min") is not None or c.get("y_max") is not None:
        (ax.set_xlim if kind == "bar" else ax.set_ylim)(c.get("y_min"), c.get("y_max"))
    if len(series) > 1 or (kind == "km" and len(c.get("arms", [])) > 1):
        ax.legend(frameon=False, loc=c.get("legend", "best"), fontsize=10)
    if c.get("title"):
        ax.set_title(c["title"], loc="left", fontweight="bold")
    fig.tight_layout()
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=dpi)
    plt.close(fig)
    return out


# ===========================================================================
# Deck: validation
# ===========================================================================

DECK_TYPES = {"title", "section", "agenda", "bullets", "two_column", "stats", "chart",
              "table", "data", "callout", "quote", "questions", "timeline", "process",
              "references", "image", "closing"}


def validate_deck(spec: dict, *, strict_approval: bool = False) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not spec.get("title"):
        errors.append("deck has no title")
    slides = spec.get("slides") or []
    if not slides:
        errors.append("deck has no slides")
    for i, s in enumerate(slides, 1):
        kind = s.get("type", "bullets")
        where = f"slide {i + 1} ({kind})"
        if kind not in DECK_TYPES:
            errors.append(f"{where}: unknown slide type — use one of {sorted(DECK_TYPES)}")
            continue
        if not s.get("title") and kind not in {"callout", "quote"}:
            errors.append(f"{where}: no title")
        if kind in {"data", "table", "chart", "stats"}:
            if not s.get("citation"):
                errors.append(f"{where}: a data slide needs a citation (source, "
                              f"identifier) on the slide itself")
            if kind in {"data", "table", "chart"} and not (s.get("design") or s.get("kicker")):
                warnings.append(f"{where}: name the study design and N in 'design' "
                                f"(shown as the kicker above the title)")
        if kind in {"data", "table"} and not s.get("rows"):
            errors.append(f"{where}: table slide has no rows")
        if kind == "chart" and not s.get("chart"):
            errors.append(f"{where}: chart slide has no 'chart' spec")
        if kind == "bullets":
            bl = s.get("bullets", [])
            if len(bl) > 6:
                warnings.append(f"{where}: {len(bl)} bullets — split the slide or move detail to notes")
            nums = re.findall(r"[−-]?\d+(?:\.\d+)?\s?%", " ".join(plain(b if isinstance(b, str) else b.get('text', '')) for b in bl))
            if len(nums) >= 3:
                warnings.append(f"{where}: {len(nums)} percentages in bullets — a 'stats' or "
                                f"'chart' slide will land these numbers better")
    if not spec.get("approval_status"):
        (errors if strict_approval else warnings).append(
            "no approval_status: state what is approved, where, and that other uses "
            "are investigational (omit only for decks that discuss no product use)")
    return errors, warnings


# ===========================================================================
# Deck: PowerPoint renderer
# ===========================================================================

class DeckBuilder:
    W, H = 13.333, 7.5
    ML = 0.65
    CW = 13.333 - 2 * 0.65
    CT = 1.75          # content top
    CB = 6.45          # content bottom

    def __init__(self, spec: dict, workdir: Path):
        from pptx import Presentation
        from pptx.util import Inches
        self.spec = spec
        self.prs = Presentation()
        self.prs.slide_width = Inches(self.W)
        self.prs.slide_height = Inches(self.H)
        self.blank = self.prs.slide_layouts[6]
        self.workdir = workdir
        self.warnings: list[str] = []
        self.page = 0
        fonts = spec.get("fonts", {})
        self.head_font = fonts.get("heading", HEAD_FONT)
        self.body_font = fonts.get("body", BODY_FONT)
        self.section_no = 0
        self.short_title = spec.get("short_title") or spec.get("title", "")[:70]
        self.charts_made = 0

    # -- primitives --------------------------------------------------------
    def rect(self, slide, x, y, w, h, fill, line=None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if line:
            shp.line.color.rgb = _rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        _plain_shape(shp)
        return shp

    def oval(self, slide, x, y, d, fill):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fill)
        shp.line.fill.background(); _plain_shape(shp)
        return shp

    def text(self, slide, x, y, w, h, paras, *, size=16, color=INK, bold=False,
             font=None, align="left", anchor="top", fit=True, min_size=None,
             bullet=None, gap=0.45, italic=False, em=None, name=None, spacing=1.12):
        """Add a text box. paras: str or list of str / {'text','sub'} items."""
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.util import Inches, Pt
        if isinstance(paras, str):
            paras = [paras]
        flat: list[tuple[str, int]] = []
        for p in paras:
            if isinstance(p, dict):
                flat.append((p.get("text", ""), 0))
                flat += [(sp, 1) for sp in p.get("sub", [])]
            else:
                flat.append((str(p), 0))
        font = font or self.body_font
        em = em or (0.6 if font == self.head_font else 0.5)
        indent_in = 0.32 if bullet else 0.0
        if fit:
            size, over = fit_size([t for t, _ in flat], w - indent_in, h, size,
                                  min_size or max(10, size - 8), em=em,
                                  spacing=spacing + 0.08, gap_ratio=gap)
            if over:
                self.warnings.append(
                    f"slide {self.page}: text may overflow ('{plain(flat[0][0])[:50]}…') — "
                    f"shorten, split the slide, or move detail to speaker notes")
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for k, (t, level) in enumerate(flat):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = spacing
            psize = size if level == 0 else max(10, size - 2)
            if k > 0:
                p.space_before = Pt(psize * gap if level == 0 else psize * 0.2)
            for (rt, rb, ri) in runs(t):
                r = p.add_run()
                r.text = rt
                r.font.size = Pt(psize)
                r.font.name = font
                r.font.bold = bold or rb
                r.font.italic = italic or ri
                r.font.color.rgb = _rgb(color if level == 0 else SLATE)
            if bullet:
                self._bullet(p, level, bullet)
        return box, size

    def _bullet(self, p, level, char):
        from pptx.oxml.ns import qn
        from lxml import etree
        pPr = p._p.get_or_add_pPr()
        mar = 320040 + level * 320040
        pPr.set("marL", str(mar))
        pPr.set("indent", str(-260000))
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont", "a:buSzPct"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        clr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(clr, qn("a:srgbClr"))
        srgb.set("val", TEAL if level == 0 else SLATE)
        sz = etree.SubElement(pPr, qn("a:buSzPct")); sz.set("val", "100000")
        fnt = etree.SubElement(pPr, qn("a:buFont")); fnt.set("typeface", "Arial")
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", ("■" if char == "square" else "•") if level == 0 else "–")

    def notes(self, slide, s):
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])

    # -- chrome --------------------------------------------------------------
    def new_slide(self):
        self.page += 1
        return self.prs.slides.add_slide(self.blank)

    def header(self, slide, s, *, top_rule=True):
        if top_rule:
            self.rect(slide, 0, 0, self.W, 0.07, TEAL)
        kicker = s.get("design") or s.get("kicker")
        y = 0.38
        if kicker:
            self.text(slide, self.ML, y, self.CW, 0.32, plain(kicker).upper(), size=11,
                      color=TEAL, bold=True, fit=True, min_size=9, name="Kicker")
        self.text(slide, self.ML, 0.68, self.CW, 0.95, s.get("title", ""), size=28,
                  font=self.head_font, bold=True, color=INK, anchor="top",
                  min_size=20, gap=0.1, name="Title")

    def footer(self, slide, s, *, dark=False):
        cite = s.get("citation", "")
        tier = s.get("tier", "")
        if tier and tier not in ("peer-reviewed", "peer reviewed"):
            cite = (cite + f"  [{tier}]").strip()
        if cite:
            self.text(slide, self.ML, 6.55, self.CW, 0.4, "Source: " + plain(cite), size=9,
                      color=SLATE, min_size=7, name="Source", gap=0.1)
        self.rect(slide, self.ML, 7.0, self.CW, 0.01, CLOUD)
        mark = DRAFT + ("   ·   " + SYNTHETIC if self.spec.get("synthetic") else "")
        self.text(slide, self.ML, 7.05, 9.6, 0.3, mark, size=8, color=OXBLOOD,
                  fit=False, name="Draft marking")
        self.text(slide, self.W - self.ML - 2.6, 7.05, 2.6, 0.3,
                  f"{self.short_title[:40]}  |  {self.page}", size=8, color=SLATE,
                  align="right", fit=False, name="Page number")

    def takeaway(self, slide, s, bottom):
        if not s.get("takeaway"):
            return bottom
        h = 0.62
        y = bottom - h
        self.rect(slide, self.ML, y, self.CW, h, TEAL_TINT)
        self.rect(slide, self.ML, y, 0.07, h, TEAL)
        self.text(slide, self.ML + 0.25, y + 0.04, self.CW - 0.4, h - 0.08, s["takeaway"],
                  size=15, bold=True, color=INK, anchor="middle", min_size=11, gap=0.1)
        return y - 0.2

    # -- slide types ---------------------------------------------------------
    def title_slide(self):
        sp = self.spec
        s = self.new_slide()
        bg = sp.get("title_image", "")
        if bg and Path(bg).exists():
            from pptx.util import Inches
            s.shapes.add_picture(str(bg), 0, 0, width=self.prs.slide_width,
                                 height=self.prs.slide_height)
            scrim = self.rect(s, 0, 0, self.W, self.H, INK)
            _alpha(scrim, 78)
        else:
            self.rect(s, 0, 0, self.W, self.H, INK)
            # quiet geometric motif, top right
            for r in range(4):
                for c in range(9):
                    self.oval(s, 9.2 + c * 0.38, 0.55 + r * 0.38, 0.07, "2A4A63")
        self.rect(s, self.ML, 2.05, 0.08, 2.35, TEAL)
        if sp.get("kicker") or sp.get("doc_type"):
            self.text(s, 0.95, 1.55, 10, 0.4, plain(sp.get("kicker") or sp.get("doc_type")).upper(),
                      size=12, bold=True, color="7FC4C2", fit=False)
        self.text(s, 0.95, 1.95, 11.3, 1.75, sp.get("title", ""), size=40,
                  font=self.head_font, bold=True, color=WHITE, min_size=24, gap=0.1,
                  anchor="bottom", name="Title", spacing=1.0)
        if sp.get("subtitle"):
            self.text(s, 0.95, 3.8, 11.3, 0.65, sp["subtitle"], size=19, color="C8D4DC",
                      min_size=13, gap=0.1)
        meta = " · ".join(str(x) for x in [sp.get("presenter", ""), sp.get("date", ""),
                                           sp.get("jurisdiction", "")] if x)
        if meta:
            self.text(s, 0.95, 4.55, 11.3, 0.4, meta, size=13, color="9FB3C2", fit=False)
        box = []
        if sp.get("approval_status"):
            box.append("**Approval status.** " + sp["approval_status"])
        if sp.get("title_note"):
            box.append(sp["title_note"])
        if box:
            self.rect(s, self.ML, 5.2, self.CW, 0.95, "1B3A52")
            self.text(s, self.ML + 0.25, 5.25, self.CW - 0.5, 0.85, box, size=12,
                      color="DDE6EC", anchor="middle", min_size=9, gap=0.2)
        mark = DRAFT + ("   ·   " + SYNTHETIC if sp.get("synthetic") else "")
        self.text(s, 0.95, 6.55, 11.8, 0.5, mark, size=11, bold=True, color="E8A0A0",
                  min_size=8, name="Draft marking")
        if sp.get("notes"):
            s.notes_slide.notes_text_frame.text = sp["notes"]

    def section(self, sd):
        s = self.new_slide()
        self.section_no += 1
        self.rect(s, 0, 0, 4.3, self.H, INK)
        self.text(s, 0.6, 2.35, 3.3, 1.6, sd.get("number", f"{self.section_no:02d}"),
                  size=88, font=self.head_font, bold=True, color="7FC4C2", fit=False)
        self.rect(s, 4.95, 2.75, 0.9, 0.06, TEAL)
        self.text(s, 4.95, 2.95, 7.7, 1.5, sd.get("title", ""), size=34,
                  font=self.head_font, bold=True, color=INK, min_size=24, gap=0.1)
        if sd.get("subtitle"):
            self.text(s, 4.95, 4.45, 7.7, 1.2, sd["subtitle"], size=17, color=SLATE, min_size=12)
        self.text(s, 4.95, 7.05, 7.7, 0.3, DRAFT, size=8, color=OXBLOOD, fit=False, name="Draft marking")
        self.text(s, self.W - self.ML - 1.0, 7.05, 1.0, 0.3, str(self.page), size=8, color=SLATE,
                  align="right", fit=False, name="Page number")
        self.notes(s, sd)

    def agenda(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        items = sd.get("items") or sd.get("bullets") or []
        n = max(1, len(items))
        top, bottom = self.CT + 0.1, self.takeaway(s, sd, self.CB)
        row_h = min(0.95, (bottom - top) / n)
        for k, it in enumerate(items):
            y = top + k * row_h
            current = sd.get("current") is not None and int(sd["current"]) == k + 1
            self.oval(s, self.ML + 0.1, y + (row_h - 0.5) / 2, 0.5, TEAL if current or sd.get("current") is None else CLOUD)
            self.text(s, self.ML + 0.1, y + (row_h - 0.5) / 2, 0.5, 0.5, str(k + 1), size=15,
                      bold=True, color=WHITE if current or sd.get("current") is None else SLATE,
                      align="center", anchor="middle", fit=False)
            self.text(s, self.ML + 0.85, y, self.CW - 1.0, row_h, it, size=20,
                      color=INK if (current or sd.get("current") is None) else SLATE,
                      anchor="middle", min_size=13, gap=0.1)
            if k < n - 1:
                self.rect(s, self.ML + 0.85, y + row_h - 0.01, self.CW - 1.0, 0.01, CLOUD)
        self.footer(s, sd)
        self.notes(s, sd)

    def bullets(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        bottom = self.takeaway(s, sd, self.CB)
        items = sd.get("bullets", [])
        side = sd.get("chart") or sd.get("image") or sd.get("stats")
        width = self.CW if not side else self.CW * 0.46
        big = sum(len(plain(b if isinstance(b, str) else b.get("text", ""))) for b in items) < 160
        self.text(s, self.ML, self.CT, width, bottom - self.CT, items, size=24 if big else 20,
                  bullet="square", min_size=12, gap=0.7)
        if side:
            x = self.ML + width + 0.35
            self._side_visual(s, sd, x, self.CT, self.CW - width - 0.35, bottom - self.CT)
        self.footer(s, sd)
        self.notes(s, sd)

    def _side_visual(self, s, sd, x, y, w, h):
        if sd.get("chart"):
            self._chart(s, sd["chart"], x, y, w, h)
        elif sd.get("image"):
            self._picture(s, sd["image"], x, y, w, h)
        elif sd.get("stats"):
            self._stat_cards(s, sd["stats"], x, y, w, h, vertical=True)

    def two_column(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        bottom = self.takeaway(s, sd, self.CB)
        gap = 0.45
        colw = (self.CW - gap) / 2
        for k, key in enumerate(("left", "right")):
            col = sd.get(key, {})
            if isinstance(col, list):
                col = {"bullets": col}
            x = self.ML + k * (colw + gap)
            accent = [INK, TEAL][k] if not col.get("style") else {
                "safety": OXBLOOD, "limitation": GOLD, "muted": SLATE}.get(col["style"], TEAL)
            y = self.CT
            if col.get("heading"):
                self.rect(s, x, y, colw, 0.55, accent)
                self.text(s, x + 0.15, y, colw - 0.3, 0.55, col["heading"], size=16, bold=True,
                          color=WHITE, anchor="middle", min_size=11, gap=0.1)
                y += 0.7
            else:
                self.rect(s, x, y, colw, 0.06, accent)
                y += 0.25
            content = col.get("bullets") or ([col["text"]] if col.get("text") else [])
            flat = [c if isinstance(c, str) else c.get("text", "") for c in content]
            need = text_height(flat, colw - 0.65, 18, spacing=1.25, gap_pt=10) + 0.5
            box_bottom = min(bottom, max(y + 2.2, y + need))
            self.rect(s, x, y - 0.1, colw, box_bottom - y + 0.1, MIST)
            self.text(s, x + 0.15, y, colw - 0.3, box_bottom - y - 0.1, content, size=18,
                      bullet="square" if col.get("bullets") else None, min_size=11, gap=0.55)
        self.footer(s, sd)
        self.notes(s, sd)

    def _stat_cards(self, s, items, x, y, w, h, vertical=False):
        n = max(1, len(items))
        gap = 0.3
        if vertical:
            ch = (h - gap * (n - 1)) / n
            boxes = [(x, y + k * (ch + gap), w, ch) for k in range(n)]
        else:
            cw = (w - gap * (n - 1)) / n
            boxes = [(x + k * (cw + gap), y, cw, h) for k in range(n)]
        for it, (bx, by, bw, bh) in zip(items, boxes):
            style = it.get("style", "")
            accent = OXBLOOD if style == "safety" else (GOLD if style == "limitation" else TEAL)
            self.rect(s, bx, by, bw, bh, MIST)
            self.rect(s, bx, by, bw, 0.07, accent)
            vh = min(1.35, bh * 0.42)
            self.text(s, bx + 0.2, by + 0.2, bw - 0.4, vh, str(it.get("value", "")),
                      size=48 if not vertical else 36, font=self.head_font, bold=True,
                      color=accent, anchor="bottom", min_size=20, gap=0.1)
            ly = by + 0.25 + vh
            self.text(s, bx + 0.2, ly, bw - 0.4, 0.75, it.get("label", ""), size=17, bold=True,
                      color=INK, min_size=11, gap=0.1)
            if it.get("note"):
                self.text(s, bx + 0.2, ly + 0.72, bw - 0.4, max(0.3, by + bh - ly - 0.8),
                          it["note"], size=13, color=SLATE, min_size=9, gap=0.2)

    def stats(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        bottom = self.takeaway(s, sd, self.CB)
        items = sd.get("items") or sd.get("stats") or []
        if len(items) > 4:
            self.warnings.append(f"slide {self.page}: {len(items)} stat cards — 4 is the maximum that reads")
        extra = sd.get("bullets")
        h = bottom - self.CT
        if extra:
            h = h * 0.62
        else:
            h = min(h, 3.3)
        self._stat_cards(s, items[:4], self.ML, self.CT, self.CW, h)
        if extra:
            self.text(s, self.ML, self.CT + h + 0.25, self.CW, bottom - self.CT - h - 0.25,
                      extra, size=16, bullet="square", min_size=11, gap=0.3)
        self.footer(s, sd)
        self.notes(s, sd)

    def _chart(self, s, chart, x, y, w, h):
        from pptx.util import Inches
        c = norm_chart(chart)
        has_ci = any(sr.get("ci_low") for sr in c["series"])
        if c["kind"] in NATIVE_KINDS and not has_ci and c.get("render") != "image":
            try:
                self._native_chart(s, c, x, y, w, h)
                return
            except Exception as exc:  # fall back to an image, never to nothing
                self.warnings.append(f"slide {self.page}: native chart failed ({exc}); used image")
        if not have("matplotlib"):
            self._table_from_chart(s, c, x, y, w, h)
            self.warnings.append(f"slide {self.page}: matplotlib unavailable — chart shown as a table")
            return
        self.charts_made += 1
        png = self.workdir / f"chart_{self.page}_{self.charts_made}.png"
        chart_png(c, png, width_in=w, height_in=h)
        s.shapes.add_picture(str(png), Inches(x), Inches(y), width=Inches(w), height=Inches(h))

    def _native_chart(self, s, c, x, y, w, h):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION,
                                     XL_LABEL_POSITION, XL_TICK_LABEL_POSITION,
                                     XL_TICK_MARK)
        from pptx.util import Inches, Pt
        cd = CategoryChartData()
        cd.categories = [str(x) for x in c["categories"]]
        for sr in c["series"]:
            cd.add_series(sr.get("name", ""), [float(v) for v in sr["values"]],
                          number_format=number_format(c))
        ctype = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "stacked": XL_CHART_TYPE.COLUMN_STACKED,
            "stacked100": XL_CHART_TYPE.COLUMN_STACKED_100,
        }[c["kind"]]
        gf = s.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.has_title = False
        ch.font.size = Pt(13)
        ch.font.name = self.body_font
        ch.font.color.rgb = _rgb(SLATE)
        multi = len(c["series"]) > 1
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(13)
        plot = ch.plots[0]
        if c["kind"] in ("column", "bar"):
            plot.gap_width = 55 if multi else 70
            plot.overlap = -8 if multi else 0
        if c.get("labels", True):
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = number_format(c)
            dl.number_format_is_linked = False
            dl.font.size = Pt(14 if not multi else 12)
            dl.font.bold = True
            dl.font.color.rgb = _rgb(INK)
            if c["kind"] in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            elif c["kind"] == "line":
                dl.position = XL_LABEL_POSITION.ABOVE
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _rgb(CLOUD)
        va.major_gridlines.format.line.width = Pt(0.75)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(11)
        va.major_tick_mark = XL_TICK_MARK.NONE
        va.tick_labels.number_format = number_format(dict(c, decimals=0))
        va.tick_labels.number_format_is_linked = False
        if c.get("y_min") is not None:
            va.minimum_scale = float(c["y_min"])
        if c.get("y_max") is not None:
            va.maximum_scale = float(c["y_max"])
        if c.get("y_label"):
            va.has_title = True
            va.axis_title.text_frame.text = c["y_label"]
            va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
        ca = ch.category_axis
        ca.format.line.color.rgb = _rgb(SLATE)
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.tick_labels.font.size = Pt(13)
        ca.tick_labels.font.color.rgb = _rgb(INK)
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        ca.has_major_gridlines = False
        hl = c.get("highlight")
        for si, series in enumerate(plot.series):
            if c["kind"] == "line":
                col = ARMS[si % len(ARMS)]
                series.format.line.color.rgb = _rgb(col)
                series.format.line.width = Pt(2.75)
                series.smooth = False
                from pptx.enum.chart import XL_MARKER_STYLE
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.size = 8
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = _rgb(col)
                series.marker.format.line.color.rgb = _rgb(col)
                continue
            try:
                series.invert_if_negative = False
            except Exception:
                pass
            fill = series.format.fill
            fill.solid()
            if multi:
                fill.fore_color.rgb = _rgb(ARMS[si % len(ARMS)])
            else:
                fill.fore_color.rgb = _rgb(INK if hl is None else MUTED_BAR)
                if hl is not None:
                    for k in range(len(c["categories"])):
                        pt = series.points[k]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = _rgb(TEAL if k == hl else MUTED_BAR)
        if c["kind"] == "bar":
            ca.reverse_order = True
            try:
                from pptx.enum.chart import XL_AXIS_CROSSES
                va.crosses = XL_AXIS_CROSSES.MAXIMUM
            except Exception:
                pass

    def _table_from_chart(self, s, c, x, y, w, h):
        rows = [[""] + [sr.get("name", "") for sr in c["series"]]]
        for k, cat in enumerate(c["categories"]):
            rows.append([cat] + [fmt_value(float(sr["values"][k]), c) for sr in c["series"]])
        self._table(s, rows, x, y, w, h, None)

    def chart(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        bottom = self.takeaway(s, sd, self.CB)
        panel = sd.get("bullets") or sd.get("insights")
        h = bottom - self.CT
        if panel:
            cw = self.CW * 0.63
            self._chart(s, sd["chart"], self.ML, self.CT, cw, h)
            px = self.ML + cw + 0.35
            pw = self.CW - cw - 0.35
            self.rect(s, px, self.CT, pw, h, MIST)
            self.rect(s, px, self.CT, pw, 0.06, TEAL)
            self.text(s, px + 0.2, self.CT + 0.2, pw - 0.4, 0.4, sd.get("panel_title", "WHAT THIS SHOWS"),
                      size=11, bold=True, color=TEAL, fit=False)
            self.text(s, px + 0.2, self.CT + 0.65, pw - 0.4, h - 0.85, panel, size=16,
                      bullet="square", min_size=10, gap=0.6)
        else:
            self._chart(s, sd["chart"], self.ML, self.CT, self.CW, h)
        self.footer(s, sd)
        self.notes(s, sd)

    def _table(self, s, rows, x, y, w, h, highlight, font_size=None):
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.oxml.ns import qn
        from lxml import etree
        n_rows, n_cols = len(rows), max(len(r) for r in rows)
        rows = [list(r) + [""] * (n_cols - len(r)) for r in rows]
        size = font_size or (16 if n_rows <= 5 else 14 if n_rows <= 7 else 12 if n_rows <= 10 else 10)
        if n_rows > 12:
            self.warnings.append(f"slide {self.page}: {n_rows}-row table — move detail to backup")
        lens = [max(len(plain(str(r[c]))) for r in rows) for c in range(n_cols)]
        weights = [min(max(L, 6), 60) for L in lens]
        tot = sum(weights)
        widths = [w * wt / tot for wt in weights]
        # row heights from wrapped text
        heights = []
        for r in rows:
            lines = max(wrapped_lines(plain(str(r[c])), max(4, int((widths[c] - 0.2) * 72 / (size * 0.5))))
                        for c in range(n_cols))
            heights.append(max(0.42, lines * size * 1.25 / 72 + 0.16))
        total_h = sum(heights)
        if total_h > h:
            self.warnings.append(f"slide {self.page}: table taller than the slide area — "
                                 f"shorten cells or split the table")
        centre_cols = [all(len(plain(str(r[c]))) <= 14 for r in rows[1:]) for c in range(n_cols)]
        gf = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(min(total_h, h)))
        tbl = gf.table
        tblPr = gf._element.graphic.graphicData.tbl.tblPr
        style_id = tblPr.find(qn("a:tableStyleId"))
        if style_id is None:
            style_id = etree.SubElement(tblPr, qn("a:tableStyleId"))
        style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
        tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
        for c in range(n_cols):
            tbl.columns[c].width = Inches(widths[c])
        for r in range(n_rows):
            tbl.rows[r].height = Inches(heights[r])
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.1)
                cell.margin_top = cell.margin_bottom = Inches(0.05)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                is_hl = highlight is not None and r == int(highlight)
                if r == 0:
                    cell.fill.fore_color.rgb = _rgb(INK)
                elif is_hl:
                    cell.fill.fore_color.rgb = _rgb(TEAL_TINT)
                else:
                    cell.fill.fore_color.rgb = _rgb(WHITE if r % 2 else MIST)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = ""
                for (rt, rb, ri) in runs(str(rows[r][c])):
                    run = p.add_run()
                    run.text = rt
                    run.font.size = Pt(size)
                    run.font.name = self.body_font
                    run.font.bold = r == 0 or rb or (c == 0) or is_hl
                    run.font.italic = ri
                    run.font.color.rgb = _rgb(WHITE if r == 0 else INK)
                if c > 0 and centre_cols[c]:
                    p.alignment = PP_ALIGN.CENTER
                # hairline bottom border
                tcPr = cell._tc.get_or_add_tcPr()
                ln = etree.SubElement(tcPr, qn("a:lnB"))
                ln.set("w", "9525")
                sf = etree.SubElement(ln, qn("a:solidFill"))
                clr = etree.SubElement(sf, qn("a:srgbClr")); clr.set("val", "C9D1D8")
        return gf

    def table(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        bottom = self.takeaway(s, sd, self.CB)
        self._table(s, sd["rows"], self.ML, self.CT, self.CW, bottom - self.CT,
                    sd.get("highlight_row"), sd.get("font_size"))
        self.footer(s, sd)
        self.notes(s, sd)

    def callout(self, sd):
        s = self.new_slide()
        style = sd.get("style", "finding")
        tint, accent, _ = CALLOUT_STYLES.get(style, CALLOUT_STYLES["finding"])
        self.rect(s, 0, 0, self.W, 0.07, accent)
        if sd.get("title"):
            self.text(s, self.ML, 0.5, self.CW, 0.4, plain(sd["title"]).upper(), size=13,
                      bold=True, color=accent, fit=False)
        self.text(s, self.ML + 0.1, 1.05, 1.0, 1.2, "“", size=110, font=self.head_font,
                  color=accent, fit=False)
        stmt = sd.get("statement") or sd.get("text") or ""
        self.text(s, 1.55, 1.55, 10.9, 3.6, stmt, size=34, font=self.head_font, bold=False,
                  color=INK, anchor="middle", min_size=20, gap=0.2)
        if sd.get("attribution"):
            self.rect(s, 1.55, 5.4, 0.8, 0.04, accent)
            self.text(s, 1.55, 5.55, 10.9, 0.6, sd["attribution"], size=15, color=SLATE, min_size=11)
        self.footer(s, sd)
        self.notes(s, sd)

    def questions(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        qs = sd.get("questions", [])
        bottom = self.takeaway(s, sd, self.CB)
        n = max(1, len(qs))
        gap = 0.18
        row_h = min(1.15, (bottom - self.CT - gap * (n - 1)) / n)
        for k, q in enumerate(qs):
            y = self.CT + k * (row_h + gap)
            self.rect(s, self.ML, y, self.CW, row_h, MIST)
            self.rect(s, self.ML, y, 0.75, row_h, TEAL)
            self.text(s, self.ML, y, 0.75, row_h, f"Q{k + 1}", size=18, bold=True, color=WHITE,
                      align="center", anchor="middle", fit=False)
            self.text(s, self.ML + 0.95, y, self.CW - 1.1, row_h, q, size=20, color=INK,
                      anchor="middle", min_size=11, gap=0.1)
        self.footer(s, sd)
        self.notes(s, sd)

    def timeline(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        items = sd.get("items", [])
        bottom = self.takeaway(s, sd, self.CB)
        n = max(1, len(items))
        mid = self.CT + (bottom - self.CT) / 2
        self.rect(s, self.ML, mid - 0.02, self.CW, 0.04, SLATE)
        slot = self.CW / n
        for k, it in enumerate(items):
            cx = self.ML + slot * (k + 0.5)
            done = it.get("status") == "done"
            self.oval(s, cx - 0.14, mid - 0.14, 0.28, TEAL if not done else SLATE)
            up = k % 2 == 0
            ty = mid - 1.75 if up else mid + 0.35
            self.text(s, cx - slot / 2 + 0.08, ty, slot - 0.16, 0.4, it.get("when", ""), size=14,
                      bold=True, color=TEAL, align="center", fit=True, min_size=10, gap=0.1)
            self.text(s, cx - slot / 2 + 0.08, ty + 0.4, slot - 0.16, 1.0, it.get("what", ""),
                      size=14, color=INK, align="center", min_size=9, gap=0.1)
        self.footer(s, sd)
        self.notes(s, sd)

    def process(self, sd):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        s = self.new_slide()
        self.header(s, sd)
        steps = sd.get("steps") or sd.get("items") or []
        bottom = self.takeaway(s, sd, self.CB)
        n = max(1, len(steps))
        gap = 0.15
        w = (self.CW - gap * (n - 1)) / n
        for k, st in enumerate(steps):
            x = self.ML + k * (w + gap)
            shp = s.shapes.add_shape(MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON,
                                     Inches(x), Inches(self.CT), Inches(w), Inches(0.8))
            shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(TEAL if k == sd.get("highlight", -1) else INK)
            shp.line.fill.background(); _plain_shape(shp)
            title = st.get("title", "") if isinstance(st, dict) else str(st)
            self.text(s, x + 0.25, self.CT, w - 0.5, 0.8, title, size=15, bold=True, color=WHITE,
                      align="center", anchor="middle", min_size=10, gap=0.1)
            if isinstance(st, dict) and st.get("text"):
                self.text(s, x + 0.05, self.CT + 1.0, w - 0.1, bottom - self.CT - 1.0,
                          st["text"] if isinstance(st["text"], list) else [st["text"]],
                          size=15, color=INK, min_size=10, gap=0.4)
        self.footer(s, sd)
        self.notes(s, sd)

    def references(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        refs = [f"{k + 1}. {r}" for k, r in enumerate(sd.get("references", []))]
        if not refs:
            refs = ["No references listed. Every data slide must cite its source; collect them here."]
        if len(refs) > 8:
            half = math.ceil(len(refs) / 2)
            colw = (self.CW - 0.4) / 2
            self.text(s, self.ML, self.CT, colw, self.CB - self.CT, refs[:half], size=12,
                      color=INK, min_size=8, gap=0.35)
            self.text(s, self.ML + colw + 0.4, self.CT, colw, self.CB - self.CT, refs[half:],
                      size=12, color=INK, min_size=8, gap=0.35)
        else:
            self.text(s, self.ML, self.CT, self.CW, self.CB - self.CT, refs, size=14,
                      color=INK, min_size=8, gap=0.45)
        self.footer(s, {})
        self.notes(s, sd)

    def _picture(self, s, path, x, y, w, h):
        from pptx.util import Inches
        p = Path(path)
        if not p.exists():
            self.text(s, x, y, w, 0.6, f"[missing image: {p}]", size=14, color=OXBLOOD, fit=False)
            self.warnings.append(f"slide {self.page}: missing image {p}")
            return
        try:
            from PIL import Image
            with Image.open(p) as im:
                iw, ih = im.size
            ratio = iw / ih
        except Exception:
            ratio = 16 / 9
        if w / h > ratio:
            ph, pw = h, h * ratio
        else:
            pw, ph = w, w / ratio
        s.shapes.add_picture(str(p), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                             width=Inches(pw), height=Inches(ph))

    def image(self, sd):
        s = self.new_slide()
        self.header(s, sd)
        bottom = self.takeaway(s, sd, self.CB)
        if sd.get("bullets"):
            iw = self.CW * 0.6
            self._picture(s, sd.get("path", ""), self.ML, self.CT, iw, bottom - self.CT)
            self.text(s, self.ML + iw + 0.35, self.CT, self.CW - iw - 0.35, bottom - self.CT,
                      sd["bullets"], size=17, bullet="square", min_size=10, gap=0.6)
        else:
            self._picture(s, sd.get("path", ""), self.ML, self.CT, self.CW, bottom - self.CT)
        self.footer(s, sd)
        self.notes(s, sd)

    def closing(self, sd):
        s = self.new_slide()
        self.rect(s, 0, 0, self.W, self.H, INK)
        self.rect(s, self.ML, 1.3, 0.08, 1.0, TEAL)
        self.text(s, 0.95, 1.2, 11.3, 1.2, sd.get("title", "Next steps"), size=36,
                  font=self.head_font, bold=True, color=WHITE, min_size=24, anchor="middle", gap=0.1)
        items = sd.get("bullets") or sd.get("items") or []
        if items:
            self.text(s, 0.95, 2.75, 11.3, 3.4, items, size=20, color="DDE6EC",
                      bullet="square", min_size=12, gap=0.6)
        self.text(s, 0.95, 6.75, 11.8, 0.4, DRAFT, size=10, bold=True, color="E8A0A0", fit=False)
        self.notes(s, sd)

    # -- orchestration ---------------------------------------------------------
    def build(self, out: Path) -> list[str]:
        self.title_slide()
        dispatch = {
            "section": self.section, "agenda": self.agenda, "bullets": self.bullets,
            "two_column": self.two_column, "stats": self.stats, "chart": self.chart,
            "table": self.table, "data": self.table, "callout": self.callout,
            "quote": self.callout, "questions": self.questions, "timeline": self.timeline,
            "process": self.process, "references": self.references, "image": self.image,
            "closing": self.closing,
        }
        for sd in self.spec.get("slides", []):
            kind = sd.get("type", "bullets")
            if kind == "title":
                continue
            dispatch[kind](sd)
        props = self.prs.core_properties
        props.title = self.spec.get("title", "")
        props.subject = self.spec.get("subtitle", "")
        props.author = self.spec.get("presenter", "Medical Affairs")
        props.comments = DRAFT
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return self.warnings


# ===========================================================================
# HTML renderers (the floor when python packages cannot be installed)
# ===========================================================================

_CSS_BASE = f"""
:root{{--ink:#{INK};--slate:#{SLATE};--teal:#{TEAL};--gold:#{GOLD};--ox:#{OXBLOOD};
--cloud:#{CLOUD};--mist:#{MIST}}}
*{{box-sizing:border-box}}
body{{margin:0;font-family:Calibri,Carlito,'Segoe UI',Arial,sans-serif;color:var(--ink);background:#fff}}
h1,h2,h3{{font-family:Georgia,'Times New Roman',serif;color:var(--ink)}}
table{{border-collapse:collapse;width:100%;margin:10px 0;font-size:.92em}}
th{{background:var(--ink);color:#fff;text-align:left;padding:7px 9px}}
td{{padding:6px 9px;border-bottom:1px solid #C9D1D8}}
tr:nth-child(even) td{{background:var(--mist)}}
.draft{{color:var(--ox);font-weight:700;font-size:.8em;letter-spacing:.02em}}
.src{{color:var(--slate);font-size:.8em}}
.callout{{padding:12px 16px;margin:14px 0;border-left:5px solid var(--teal);background:#{TEAL_TINT}}}
.callout.safety{{border-color:var(--ox);background:#{OX_TINT}}}
.callout.limitation{{border-color:var(--gold);background:#{GOLD_TINT}}}
.callout.note{{border-color:var(--slate);background:var(--mist)}}
.callout b.t{{display:block;font-size:.75em;letter-spacing:.08em;color:var(--slate)}}
.kpis{{display:flex;gap:14px}} .kpi{{flex:1;background:var(--mist);border-top:4px solid var(--teal);padding:12px}}
.kpi .v{{font:700 2em Georgia,serif;color:var(--teal)}}
"""


def _h(text) -> str:
    out = []
    for t, b, i in runs(str(text)):
        t = html.escape(t)
        if b:
            t = f"<b>{t}</b>"
        if i:
            t = f"<i>{t}</i>"
        out.append(t)
    return "".join(out)


def deck_html(spec: dict) -> str:
    css = _CSS_BASE + """
@page{size:13.333in 7.5in;margin:0}
.slide{width:13.333in;height:7.5in;position:relative;overflow:hidden;page-break-after:always;
padding:.45in .65in;border-bottom:1px solid #ddd}
.slide .k{color:var(--teal);font-weight:700;font-size:11pt;letter-spacing:.06em;text-transform:uppercase}
.slide h2{font-size:26pt;margin:.1in 0 .25in}
.slide li{font-size:18pt;margin:.1in 0}
.slide .ft{position:absolute;left:.65in;right:.65in;bottom:.2in;display:flex;justify-content:space-between;
font-size:8pt;color:var(--slate);border-top:1px solid var(--cloud);padding-top:4px}
.slide .cite{position:absolute;left:.65in;right:.65in;bottom:.6in;font-size:9pt;color:var(--slate)}
.title{background:var(--ink);color:#fff}.title h1{color:#fff;font-size:40pt;margin-top:1.4in}
.title .sub{color:#C8D4DC;font-size:19pt}.title .ap{background:#1B3A52;padding:10px 14px;margin-top:.4in;font-size:12pt}
"""
    parts = [f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(spec.get('title', ''))}</title>"
             f"<style>{css}</style></head><body>"]
    parts.append(f"<section class='slide title'><h1>{_h(spec.get('title', ''))}</h1>"
                 f"<div class='sub'>{_h(spec.get('subtitle', ''))}</div>"
                 f"<p>{html.escape(' · '.join(str(x) for x in [spec.get('presenter', ''), spec.get('date', '')] if x))}</p>"
                 + (f"<div class='ap'><b>Approval status.</b> {_h(spec['approval_status'])}</div>" if spec.get('approval_status') else "")
                 + f"<p class='draft' style='color:#E8A0A0'>{DRAFT}</p></section>")
    for k, sd in enumerate(spec.get("slides", []), 2):
        kind = sd.get("type", "bullets")
        body = []
        if sd.get("design") or sd.get("kicker"):
            body.append(f"<div class='k'>{_h(sd.get('design') or sd.get('kicker'))}</div>")
        body.append(f"<h2>{_h(sd.get('title', sd.get('statement', '')))}</h2>")
        items = (sd.get("bullets") or sd.get("items") or sd.get("questions")
                 or sd.get("references") or [])
        if kind in ("data", "table") and sd.get("rows"):
            rows = sd["rows"]
            body.append("<table><tr>" + "".join(f"<th>{_h(c)}</th>" for c in rows[0]) + "</tr>"
                        + "".join("<tr>" + "".join(f"<td>{_h(c)}</td>" for c in r) + "</tr>" for r in rows[1:])
                        + "</table>")
        elif kind == "stats":
            body.append("<div class='kpis'>" + "".join(
                f"<div class='kpi'><div class='v'>{_h(i.get('value', ''))}</div><b>{_h(i.get('label', ''))}</b>"
                f"<div class='src'>{_h(i.get('note', ''))}</div></div>" for i in (sd.get('items') or [])) + "</div>")
        elif kind == "chart":
            c = norm_chart(sd["chart"])
            rows = [[""] + [s.get("name", "") for s in c["series"]]] + [
                [cat] + [fmt_value(float(s["values"][j]), c) for s in c["series"]]
                for j, cat in enumerate(c["categories"])]
            if c["kind"] == "forest":
                rows = [["", "Estimate (95% CI)"]] + [[r["label"], f"{r['est']} ({r['lo']}–{r['hi']})"] for r in c.get("rows", [])]
            body.append("<table><tr>" + "".join(f"<th>{_h(x)}</th>" for x in rows[0]) + "</tr>"
                        + "".join("<tr>" + "".join(f"<td>{_h(x)}</td>" for x in r) + "</tr>" for r in rows[1:])
                        + "</table>")
        elif kind == "two_column":
            for key in ("left", "right"):
                col = sd.get(key, {})
                if isinstance(col, list):
                    col = {"bullets": col}
                body.append(f"<div style='display:inline-block;width:48%;vertical-align:top'><h3>{_h(col.get('heading', ''))}</h3><ul>"
                            + "".join(f"<li>{_h(b)}</li>" for b in col.get("bullets", [])) + "</ul></div>")
        elif kind in ("callout", "quote"):
            body.append(f"<p style='font:28pt Georgia,serif'>{_h(sd.get('statement') or sd.get('text', ''))}</p>")
        if kind == "timeline":
            items = [f"**{i.get('when', '')}** — {i.get('what', '')}" for i in items]
        if kind == "process":
            items = [(i.get("title", "") + (": " + str(i.get("text", "")) if i.get("text") else "")) if isinstance(i, dict) else i for i in (sd.get("steps") or items)]
        if items and kind not in ("stats",):
            tag = "ol" if kind in ("questions", "references", "agenda") else "ul"
            body.append(f"<{tag}>" + "".join(
                f"<li>{_h(i if isinstance(i, str) else i.get('text', ''))}</li>" for i in items) + f"</{tag}>")
        if sd.get("takeaway"):
            body.append(f"<div class='callout'><b>{_h(sd['takeaway'])}</b></div>")
        if sd.get("citation"):
            body.append(f"<div class='cite'>Source: {_h(sd['citation'])}</div>")
        body.append(f"<div class='ft'><span class='draft'>{DRAFT}</span><span>{k}</span></div>")
        if sd.get("notes"):
            body.append(f"<!-- speaker notes: {html.escape(str(sd['notes']))} -->")
        parts.append("<section class='slide'>" + "".join(body) + "</section>")
    parts.append("</body></html>")
    return "".join(parts)


def report_html(spec: dict) -> str:
    css = _CSS_BASE + """
@page{size:A4;margin:22mm 18mm;@bottom-center{content:"Page " counter(page) " of " counter(pages);font-size:8pt;color:#5B6B79}}
.wrap{max-width:820px;margin:0 auto;padding:36px 24px;line-height:1.5}
.cover{border-top:8px solid var(--teal);padding-top:20px;margin-bottom:26px}
.cover .k{color:var(--teal);font-weight:700;letter-spacing:.08em;text-transform:uppercase;font-size:.8em}
.cover h1{font-size:2.2em;margin:.2em 0}
h2{border-bottom:2px solid var(--teal);padding-bottom:4px;margin-top:1.6em}
.summary{background:var(--mist);border-left:5px solid var(--ink);padding:12px 18px}
"""
    p = [f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(spec.get('title', ''))}</title>"
         f"<style>{css}</style></head><body><div class='wrap'>"]
    p.append(f"<div class='draft'>{DRAFT}</div><div class='cover'><div class='k'>{_h(spec.get('doc_type', 'Medical Affairs report'))}</div>"
             f"<h1>{_h(spec.get('title', ''))}</h1><div>{_h(spec.get('subtitle', ''))}</div>"
             f"<div class='src'>{html.escape(' · '.join(str(x) for x in [spec.get('author', ''), spec.get('date', ''), spec.get('doc_id', ''), ('v' + str(spec['version'])) if spec.get('version') else ''] if x))}</div></div>")
    if spec.get("approval_status"):
        p.append(f"<div class='callout note'><b class='t'>APPROVAL STATUS</b>{_h(spec['approval_status'])}</div>")
    if spec.get("summary"):
        p.append("<div class='summary'><h3>Key takeaways</h3><ol>" + "".join(
            f"<li>{_h(x)}</li>" for x in spec["summary"]) + "</ol></div>")
    for sec in spec.get("sections", []):
        if sec.get("heading"):
            p.append(f"<h2>{_h(sec['heading'])}</h2>")
        for b in _section_blocks(sec):
            p.append(_block_html(b))
    if spec.get("references"):
        p.append("<h2>References</h2><ol class='src'>" + "".join(
            f"<li>{_h(r)}</li>" for r in spec["references"]) + "</ol>")
    p.append(f"<p class='draft'>{DRAFT}</p></div></body></html>")
    return "".join(p)


def _block_html(b: dict) -> str:
    t = b.get("type", "paragraph")
    if t == "paragraph":
        return f"<p>{_h(b.get('text', ''))}</p>"
    if t == "heading":
        lv = min(4, int(b.get("level", 2)) + 1)
        return f"<h{lv}>{_h(b.get('text', ''))}</h{lv}>"
    if t in ("bullets", "numbered"):
        tag = "ol" if t == "numbered" else "ul"
        items = []
        for it in b.get("items", []):
            if isinstance(it, dict):
                items.append(f"<li>{_h(it.get('text', ''))}<ul>" + "".join(
                    f"<li>{_h(x)}</li>" for x in it.get("sub", [])) + "</ul></li>")
            else:
                items.append(f"<li>{_h(it)}</li>")
        return f"<{tag}>" + "".join(items) + f"</{tag}>"
    if t == "table":
        rows = b.get("rows", [])
        if not rows:
            return ""
        cap = f"<p><b>{_h(b['caption'])}</b></p>" if b.get("caption") else ""
        src = f"<p class='src'>Source: {_h(b['source'])}</p>" if b.get("source") else ""
        return (cap + "<table><tr>" + "".join(f"<th>{_h(c)}</th>" for c in rows[0]) + "</tr>"
                + "".join("<tr>" + "".join(f"<td>{_h(c)}</td>" for c in r) + "</tr>" for r in rows[1:])
                + "</table>" + src)
    if t == "callout":
        style = b.get("style", "note")
        label = b.get("title") or CALLOUT_STYLES.get(style, CALLOUT_STYLES["note"])[2]
        return f"<div class='callout {style}'><b class='t'>{_h(label).upper()}</b>{_h(b.get('text', ''))}</div>"
    if t == "kpis":
        return "<div class='kpis'>" + "".join(
            f"<div class='kpi'><div class='v'>{_h(i.get('value', ''))}</div><b>{_h(i.get('label', ''))}</b>"
            f"<div class='src'>{_h(i.get('note', ''))}</div></div>" for i in b.get("items", [])) + "</div>"
    if t == "chart":
        c = norm_chart(b.get("chart", {}))
        if c["kind"] == "forest":
            rows = [["", "Estimate (95% CI)"]] + [[r["label"], f"{r['est']} ({r['lo']}–{r['hi']})"] for r in c.get("rows", [])]
        else:
            rows = [[""] + [s.get("name", "") for s in c["series"]]] + [
                [cat] + [fmt_value(float(s["values"][j]), c) for s in c["series"]]
                for j, cat in enumerate(c["categories"])]
        return _block_html({"type": "table", "rows": rows, "caption": b.get("caption", ""),
                            "source": b.get("source", "")})
    if t == "image":
        return f"<figure><img src='{html.escape(b.get('path', ''))}' style='max-width:100%'><figcaption class='src'>{_h(b.get('caption', ''))}</figcaption></figure>"
    if t == "quote":
        return f"<blockquote><i>{_h(b.get('text', ''))}</i><br><span class='src'>{_h(b.get('attribution', ''))}</span></blockquote>"
    if t == "source":
        return f"<p class='src'>Source: {_h(b.get('text', ''))}</p>"
    if t == "page_break":
        return "<div style='page-break-after:always'></div>"
    return f"<p>{_h(b.get('text', ''))}</p>"


def _section_blocks(sec: dict) -> list[dict]:
    blocks = [dict(b, type="kpis") if b.get("type") == "stats" else b
              for b in (sec.get("blocks") or [])]
    if sec.get("body") and not blocks:
        body = sec["body"]
        blocks = [{"type": "paragraph", "text": t} for t in (body if isinstance(body, list) else [body])]
    if sec.get("bullets"):
        blocks.append({"type": "bullets", "items": sec["bullets"]})
    return blocks


# ===========================================================================
# Report: Word renderer (python-docx)
# ===========================================================================

def validate_report(spec: dict) -> tuple[list[str], list[str]]:
    errors, warnings = [], []
    if not spec.get("title"):
        errors.append("report has no title")
    if not spec.get("sections"):
        errors.append("report has no sections")
    for i, sec in enumerate(spec.get("sections", []), 1):
        for b in sec.get("blocks", []) or []:
            if b.get("type") in ("table", "chart", "kpis") and not (b.get("source") or spec.get("references")):
                warnings.append(f"section {i} '{sec.get('heading', '')[:40]}': "
                                f"{b['type']} without a source line")
    if not spec.get("summary"):
        warnings.append("no 'summary' — a three-to-five point executive summary makes the "
                        "document usable by someone who reads one page")
    return errors, warnings


class ReportBuilder:
    def __init__(self, spec: dict, workdir: Path):
        from docx import Document
        self.spec = spec
        self.workdir = workdir
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self.warnings: list[str] = []
        fonts = spec.get("fonts", {})
        self.head_font = fonts.get("heading", HEAD_FONT)
        self.body_font = fonts.get("body", BODY_FONT)
        self._styles()

    # -- low level ------------------------------------------------------------
    def _styles(self):
        from docx.shared import Pt, RGBColor, Inches
        from docx.oxml.ns import qn
        d = self.doc
        sec = d.sections[0]
        if self.spec.get("page_size", "letter").lower() == "a4":
            sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(0.95)
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.85)
        sec.different_first_page_header_footer = True
        st = d.styles["Normal"]
        st.font.name = self.body_font
        st.font.size = Pt(10.5)
        st.font.color.rgb = RGBColor.from_string("1F2D3A")
        st.element.rPr.rFonts.set(qn("w:eastAsia"), self.body_font)
        pf = st.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        for name, size, color, before, after in (("Heading 1", 17, INK, 18, 6),
                                                 ("Heading 2", 13, TEAL, 14, 4),
                                                 ("Heading 3", 11.5, INK, 10, 3)):
            h = d.styles[name]
            h.font.name = self.head_font
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = RGBColor.from_string(color)
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                from docx.oxml import OxmlElement
                rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), self.head_font)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number", "List Bullet 2"):
            try:
                ls = d.styles[name]
                ls.font.name = self.body_font
                ls.font.size = Pt(10.5)
                ls.paragraph_format.space_after = Pt(3)
            except KeyError:
                pass

    def _shade(self, cell, fill):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def _borders(self, cell, **edges):
        """edges: top/left/bottom/right = (size_eighths, color) or None for nil."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders"); tcPr.append(b)
        for edge in ("top", "left", "bottom", "right"):
            v = edges.get(edge, "nil")
            el = OxmlElement(f"w:{edge}")
            if v == "nil" or v is None:
                el.set(qn("w:val"), "nil")
            else:
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(v[0]))
                el.set(qn("w:space"), "0"); el.set(qn("w:color"), v[1])
            b.append(el)

    def _cell_margins(self, table, top=60, bottom=60, left=110, right=110):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tblPr = table._tbl.tblPr
        m = OxmlElement("w:tblCellMar")
        for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
        tblPr.append(m)

    def _runs(self, p, text, *, size=None, bold=False, italic=False, color=None, font=None):
        from docx.shared import Pt, RGBColor
        for t, b, i in runs(str(text)):
            r = p.add_run(t)
            r.bold = bold or b or None
            r.italic = italic or i or None
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            if font:
                r.font.name = font
        return p

    def _field(self, p, instr, size=8, color=SLATE):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
        r = p.add_run()
        r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
        for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
            if kind == "instr":
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {text} "
                r._r.append(it)
            elif kind == "text":
                t = OxmlElement("w:t"); t.text = text; r._r.append(t)
            else:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)

    def _para(self, text="", *, style=None, size=None, bold=False, italic=False, color=None,
              align=None, after=None, before=None, font=None, keep=False):
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = self.doc.add_paragraph(style=style)
        if text:
            self._runs(p, text, size=size, bold=bold, italic=italic, color=color, font=font)
        if align:
            p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                           "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    # -- page furniture ---------------------------------------------------------
    def header_footer(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, Inches
        sec = self.doc.sections[0]
        sp = self.spec
        width = sec.page_width - sec.left_margin - sec.right_margin
        ht = sec.header.add_table(rows=1, cols=2, width=width)
        ht.autofit = False
        left, right = ht.cell(0, 0), ht.cell(0, 1)
        left.width, right.width = int(width * 0.68), int(width * 0.32)
        for cell in (left, right):
            self._borders(cell, bottom=(4, CLOUD))
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        self._runs(left.paragraphs[0], (sp.get("short_title") or sp.get("title", ""))[:90], size=8, color=SLATE)
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._runs(right.paragraphs[0], "DRAFT — NOT FOR EXTERNAL USE", size=8, bold=True, color=OXBLOOD)
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        ctrl = " · ".join(str(x) for x in [sp.get("doc_id", ""), ("Version " + str(sp["version"])) if sp.get("version") else "",
                                           sp.get("date", "")] if x)
        self._runs(fp, (ctrl + "    ") if ctrl else "", size=8, color=SLATE)
        self._runs(fp, "Page ", size=8, color=SLATE)
        self._field(fp, "PAGE")
        self._runs(fp, " of ", size=8, color=SLATE)
        self._field(fp, "NUMPAGES")
        fp2 = sec.footer.add_paragraph()
        self._runs(fp2, DRAFT + ("  ·  " + SYNTHETIC if sp.get("synthetic") else ""), size=7, color=OXBLOOD)
        # first page footer: draft marking only
        ffp = sec.first_page_footer.paragraphs[0]
        self._runs(ffp, DRAFT, size=8, bold=True, color=OXBLOOD)
        ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def cover(self):
        from docx.shared import Pt, Inches
        from docx.enum.table import WD_TABLE_ALIGNMENT
        sp = self.spec
        width = self.doc.sections[0].page_width - self.doc.sections[0].left_margin - self.doc.sections[0].right_margin
        band = self.doc.add_table(rows=1, cols=1)
        band.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = band.cell(0, 0)
        c.width = width
        self._shade(c, INK)
        self._borders(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(26)
        self._runs(p, plain(sp.get("doc_type", "Medical Affairs report")).upper(), size=9.5, bold=True, color="7FC4C2")
        p2 = c.add_paragraph()
        p2.paragraph_format.space_before = Pt(8)
        p2.paragraph_format.line_spacing = 1.0
        self._runs(p2, sp.get("title", ""), size=26, bold=True, color=WHITE, font=self.head_font)
        if sp.get("subtitle"):
            p3 = c.add_paragraph()
            p3.paragraph_format.space_before = Pt(6)
            self._runs(p3, sp["subtitle"], size=13, color="C8D4DC")
        p4 = c.add_paragraph()
        p4.paragraph_format.space_after = Pt(26)
        teal = self.doc.add_table(rows=1, cols=1)
        tc = teal.cell(0, 0)
        self._shade(tc, TEAL); self._borders(tc)
        tp = tc.paragraphs[0]
        tp.paragraph_format.space_after = Pt(0)
        r = tp.add_run(" "); r.font.size = Pt(3)
        self._para("", after=10)
        # document control table
        meta = [("Prepared by", sp.get("author", "")), ("Date", sp.get("date", str(date.today()))),
                ("Version", sp.get("version", "")), ("Document ID", sp.get("doc_id", "")),
                ("Audience", sp.get("audience", "")), ("Status", sp.get("status", "Draft for qualified medical review"))]
        meta = [(k, v) for k, v in meta if v]
        t = self.doc.add_table(rows=len(meta), cols=2)
        self._cell_margins(t, 50, 50, 80, 80)
        for k, (label, val) in enumerate(meta):
            a, b = t.cell(k, 0), t.cell(k, 1)
            a.width, b.width = Inches(1.6), width - Inches(1.6)
            self._borders(a, bottom=(4, "D5DCE2")); self._borders(b, bottom=(4, "D5DCE2"))
            a.paragraphs[0].paragraph_format.space_after = Pt(0)
            b.paragraphs[0].paragraph_format.space_after = Pt(0)
            self._runs(a.paragraphs[0], label.upper(), size=8, bold=True, color=SLATE)
            self._runs(b.paragraphs[0], str(val), size=10)
        self._para("", after=6)
        if sp.get("approval_status"):
            self.callout({"style": "note", "title": "Approval status", "text": sp["approval_status"]})
        mark = DRAFT + ("\n" + SYNTHETIC if sp.get("synthetic") else "")
        self.callout({"style": "safety", "title": "Document status", "text": mark})
        if sp.get("summary"):
            self.summary(sp["summary"])
        if sp.get("toc"):
            self.doc.add_page_break()
            self._para("Contents", style="Heading 1")
            p = self.doc.add_paragraph()
            self._field(p, 'TOC \\o "1-2" \\h \\z \\u', size=10, color=INK)
        self.doc.add_page_break()

    def summary(self, points):
        from docx.shared import Pt
        t = self.doc.add_table(rows=1, cols=1)
        c = t.cell(0, 0)
        self._shade(c, MIST)
        self._borders(c, left=(36, INK))
        self._cell_margins(t, 120, 120, 200, 160)
        p = c.paragraphs[0]
        self._runs(p, "KEY TAKEAWAYS", size=9, bold=True, color=TEAL)
        p.paragraph_format.space_after = Pt(4)
        for k, pt in enumerate(points, 1):
            q = c.add_paragraph()
            q.paragraph_format.left_indent = Pt(18)
            q.paragraph_format.first_line_indent = Pt(-18)
            q.paragraph_format.space_after = Pt(4)
            self._runs(q, f"{k}.\u00a0\u00a0", bold=True, color=TEAL, size=11)
            self._runs(q, pt, size=11)
        self._para("", after=4)

    # -- blocks -------------------------------------------------------------------
    def callout(self, b):
        from docx.shared import Pt
        style = b.get("style", "note")
        tint, accent, label = CALLOUT_STYLES.get(style, CALLOUT_STYLES["note"])
        t = self.doc.add_table(rows=1, cols=1)
        c = t.cell(0, 0)
        self._shade(c, tint)
        self._borders(c, left=(30, accent))
        self._cell_margins(t, 90, 90, 180, 140)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        self._runs(p, plain(b.get("title") or label).upper(), size=8.5, bold=True, color=accent)
        for k, line in enumerate(str(b.get("text", "")).split("\n")):
            q = c.add_paragraph()
            q.paragraph_format.space_after = Pt(2)
            self._runs(q, line, size=10.5)
        for item in b.get("bullets", []) or []:
            q = c.add_paragraph()
            q.paragraph_format.left_indent = Pt(12)
            self._runs(q, "•  " + str(item), size=10.5)
        self._para("", after=2)

    def table(self, b):
        from docx.shared import Pt, Inches
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        rows = b.get("rows") or []
        if not rows:
            return
        n_cols = max(len(r) for r in rows)
        rows = [list(r) + [""] * (n_cols - len(r)) for r in rows]
        if b.get("caption"):
            self.tab_no += 1
            p = self._para("", keep=True, after=3, before=6)
            self._runs(p, f"Table {self.tab_no}. ", bold=True, color=TEAL, size=9.5)
            self._runs(p, b["caption"], bold=True, size=9.5)
        t = self.doc.add_table(rows=len(rows), cols=n_cols)
        t.autofit = False
        self._cell_margins(t, 45, 45, 90, 90)
        size = 9.5 if n_cols <= 4 else 8.5
        width = self.doc.sections[0].page_width - self.doc.sections[0].left_margin - self.doc.sections[0].right_margin
        lens = [max(len(plain(str(r[c]))) for r in rows) for c in range(n_cols)]
        weights = [min(max(L, 8), 55) for L in lens]
        tot = sum(weights)
        hl = b.get("highlight_row")
        centre = [c > 0 and all(len(plain(str(r[c]))) <= 12 for r in rows[1:]) for c in range(n_cols)]
        for c in range(n_cols):
            t.columns[c].width = int(width * weights[c] / tot)
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        for r, row in enumerate(rows):
            trPr = t.rows[r]._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit"); trPr.append(cant)
            if r == 0:
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            for c, val in enumerate(row):
                cell = t.cell(r, c)
                cell.width = int(width * weights[c] / tot)
                fill = INK if r == 0 else (TEAL_TINT if hl is not None and r == int(hl) else (MIST if r % 2 == 0 else WHITE))
                self._shade(cell, fill)
                self._borders(cell, bottom=(4, "C9D1D8"))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if centre[c]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._runs(p, str(val), size=size, bold=(r == 0 or (c == 0 and b.get("bold_first_col", True))),
                           color=WHITE if r == 0 else None)
        if b.get("source") or b.get("note"):
            p = self._para("", after=8, before=3)
            if b.get("note"):
                self._runs(p, b["note"] + "  ", italic=True, size=8, color=SLATE)
            if b.get("source"):
                self._runs(p, "Source: " + b["source"], size=8, color=SLATE)
        else:
            self._para("", after=4)

    def kpis(self, b):
        from docx.shared import Pt
        items = b.get("items", [])[:4]
        if not items:
            return
        t = self.doc.add_table(rows=1, cols=len(items))
        self._cell_margins(t, 100, 100, 120, 120)
        for k, it in enumerate(items):
            c = t.cell(0, k)
            accent = OXBLOOD if it.get("style") == "safety" else TEAL
            self._shade(c, MIST)
            self._borders(c, top=(24, accent), left=None if k == 0 else (8, WHITE))
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            self._runs(p, str(it.get("value", "")), size=22, bold=True, color=accent, font=self.head_font)
            q = c.add_paragraph(); q.paragraph_format.space_after = Pt(1)
            self._runs(q, it.get("label", ""), size=9.5, bold=True)
            if it.get("note"):
                r = c.add_paragraph(); r.paragraph_format.space_after = Pt(0)
                self._runs(r, it["note"], size=8.5, color=SLATE)
        if b.get("source"):
            p = self._para("", after=8, before=3)
            self._runs(p, "Source: " + b["source"], size=8, color=SLATE)
        else:
            self._para("", after=4)

    def figure(self, path: Path, caption="", source="", width_in=6.4):
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width_in))
        if caption or source:
            self.fig_no += 1
            q = self._para("", after=8)
            if caption:
                self._runs(q, f"Figure {self.fig_no}. ", bold=True, color=TEAL, size=9)
                self._runs(q, caption + ("  " if source else ""), size=9, bold=True)
            if source:
                self._runs(q, "Source: " + source, size=8, color=SLATE)

    def chart(self, b):
        if not have("matplotlib"):
            self.warnings.append("matplotlib unavailable — chart rendered as a table")
            c = norm_chart(b.get("chart", {}))
            if c["kind"] == "forest":
                rows = [["", "Estimate (95% CI)"]] + [[r["label"], f"{r['est']} ({r['lo']}–{r['hi']})"] for r in c.get("rows", [])]
            else:
                rows = [[""] + [s.get("name", "") for s in c["series"]]] + [
                    [cat] + [fmt_value(float(s["values"][j]), c) for s in c["series"]]
                    for j, cat in enumerate(c["categories"])]
            self.table({"rows": rows, "caption": b.get("caption", ""), "source": b.get("source", "")})
            return
        self.fig_png = getattr(self, "fig_png", 0) + 1
        png = self.workdir / f"figure_{self.fig_png}.png"
        chart_png(b["chart"], png, width_in=b.get("width_in", 7.2), height_in=b.get("height_in", 3.6))
        self.figure(png, b.get("caption", ""), b.get("source", ""), width_in=min(6.5, b.get("width_in", 6.5)))

    def bullets(self, items, numbered=False):
        from docx.shared import Pt
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            text = it if isinstance(it, str) else it.get("text", "")
            p = self.doc.add_paragraph(style=style)
            self._runs(p, text)
            if isinstance(it, dict):
                for sub in it.get("sub", []):
                    q = self.doc.add_paragraph(style="List Bullet 2")
                    self._runs(q, sub, color="3A4A58")

    def block(self, b):
        from docx.shared import Pt, Inches
        t = b.get("type", "paragraph")
        if t == "paragraph":
            self._para(b.get("text", ""))
        elif t == "heading":
            lvl = int(b.get("level", 2))
            self.doc.add_heading(plain(b.get("text", "")), level=min(3, max(1, lvl)))
        elif t == "bullets":
            self.bullets(b.get("items", []))
        elif t == "numbered":
            self.bullets(b.get("items", []), numbered=True)
        elif t == "table":
            self.table(b)
        elif t == "callout":
            self.callout(b)
        elif t == "kpis":
            self.kpis(b)
        elif t == "chart":
            self.chart(b)
        elif t == "image":
            p = Path(b.get("path", ""))
            if p.exists():
                self.figure(p, b.get("caption", ""), b.get("source", ""), b.get("width_in", 6.4))
            else:
                self.warnings.append(f"missing image {p}")
                self._para(f"[Missing figure: {p}]", color=OXBLOOD)
        elif t == "quote":
            p = self._para(b.get("text", ""), italic=True, size=11.5, color=INK)
            p.paragraph_format.left_indent = Inches(0.4)
            if b.get("attribution"):
                q = self._para("— " + b["attribution"], size=9, color=SLATE)
                q.paragraph_format.left_indent = Inches(0.4)
        elif t == "source":
            self._para("Source: " + b.get("text", ""), size=8, color=SLATE)
        elif t == "page_break":
            self.doc.add_page_break()
        else:
            self._para(b.get("text", ""))

    def build(self, out: Path) -> list[str]:
        sp = self.spec
        self.header_footer()
        self.cover()
        for k, sec in enumerate(sp.get("sections", []), 1):
            if sec.get("heading"):
                label = f"{k}.  {sec['heading']}" if sp.get("numbered_sections", True) else sec["heading"]
                self.doc.add_heading(plain(label), level=1)
            for b in _section_blocks(sec):
                self.block(b)
            if sec.get("page_break"):
                self.doc.add_page_break()
        if sp.get("references"):
            self.doc.add_heading("References", level=1)
            from docx.shared import Pt
            for k, ref in enumerate(sp["references"], 1):
                p = self._para("", after=3)
                p.paragraph_format.left_indent = Pt(20)
                p.paragraph_format.first_line_indent = Pt(-20)
                self._runs(p, f"{k}.\u00a0\u00a0", size=9, bold=True, color=TEAL)
                self._runs(p, ref, size=9)
        if sp.get("appendix"):
            self.doc.add_page_break()
            self.doc.add_heading("Appendix", level=1)
            for b in sp["appendix"]:
                self.block(b)
        cp = self.doc.core_properties
        cp.title = sp.get("title", "")
        cp.subject = sp.get("subtitle", "")
        cp.author = sp.get("author", "Medical Affairs")
        cp.comments = DRAFT
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        return self.warnings


# ===========================================================================
# Report: standard-library .docx writer (no packages at all). Plainer, but a
# genuine Word document with styles, tables, callouts and the draft marking.
# ===========================================================================

def _x(t) -> str:
    return html.escape(str(t), quote=False)


def _wruns(text, size=None, color=None, bold=False, italic=False) -> str:
    out = []
    for t, b, i in runs(str(text)):
        rpr = ""
        if bold or b:
            rpr += "<w:b/>"
        if italic or i:
            rpr += "<w:i/>"
        if color:
            rpr += f'<w:color w:val="{color}"/>'
        if size:
            rpr += f'<w:sz w:val="{int(size * 2)}"/>'
        out.append(f'<w:r><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{_x(t)}</w:t></w:r>')
    return "".join(out)


def _wp(text="", style=None, **kw) -> str:
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{ppr}{_wruns(text, **kw)}</w:p>"


def _wtable(rows, header=True, fills=None) -> str:
    n = max(len(r) for r in rows)
    grid = "".join('<w:gridCol w:w="%d"/>' % (9360 // n) for _ in range(n))
    out = ['<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="pct"/><w:tblBorders>'
           '<w:bottom w:val="single" w:sz="4" w:color="C9D1D8"/>'
           '<w:insideH w:val="single" w:sz="4" w:color="C9D1D8"/></w:tblBorders>'
           '<w:tblCellMar><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar>'
           f'</w:tblPr><w:tblGrid>{grid}</w:tblGrid>']
    for r, row in enumerate(rows):
        row = list(row) + [""] * (n - len(row))
        hdr = header and r == 0
        out.append("<w:tr>" + ("<w:trPr><w:tblHeader/></w:trPr>" if hdr else ""))
        for c, val in enumerate(row):
            fill = INK if hdr else ((fills or {}).get(r) or (MIST if r % 2 == 0 else "FFFFFF"))
            out.append(f'<w:tc><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="{fill}"/></w:tcPr>'
                       f'<w:p><w:pPr><w:spacing w:after="0"/></w:pPr>'
                       f'{_wruns(val, size=9.5, color="FFFFFF" if hdr else None, bold=hdr or c == 0)}</w:p></w:tc>')
        out.append("</w:tr>")
    out.append("</w:tbl>")
    return "".join(out)


def _wcallout(title, text, style) -> str:
    tint, accent, label = CALLOUT_STYLES.get(style, CALLOUT_STYLES["note"])
    return ('<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="pct"/><w:tblBorders>'
            f'<w:left w:val="single" w:sz="30" w:color="{accent}"/></w:tblBorders></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="9360"/></w:tblGrid><w:tr><w:tc>'
            f'<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="{tint}"/></w:tcPr>'
            f'{_wp(plain(title or label).upper(), size=8.5, color=accent, bold=True)}'
            + "".join(_wp(line) for line in str(text).split("\n"))
            + '</w:tc></w:tr></w:tbl>' + _wp(""))


def _wblock(b: dict) -> str:
    t = b.get("type", "paragraph")
    if t == "paragraph":
        return _wp(b.get("text", ""))
    if t == "heading":
        return _wp(plain(b.get("text", "")), style=f"Heading{min(3, max(1, int(b.get('level', 2))))}")
    if t in ("bullets", "numbered"):
        out = []
        for k, it in enumerate(b.get("items", []), 1):
            lead = f"{k}.  " if t == "numbered" else "•  "
            text = it if isinstance(it, str) else it.get("text", "")
            out.append(_wp(lead + text, style="ListPara"))
            if isinstance(it, dict):
                out += [_wp("–  " + s, style="ListPara2") for s in it.get("sub", [])]
        return "".join(out)
    if t == "table":
        cap = _wp(b["caption"], bold=True, size=9.5) if b.get("caption") else ""
        src = _wp("Source: " + b["source"], size=8, color=SLATE) if b.get("source") else _wp("")
        return cap + _wtable(b.get("rows", [])) + src
    if t == "callout":
        return _wcallout(b.get("title", ""), b.get("text", ""), b.get("style", "note"))
    if t == "kpis":
        rows = [[i.get("label", "") for i in b.get("items", [])],
                [str(i.get("value", "")) for i in b.get("items", [])],
                [i.get("note", "") for i in b.get("items", [])]]
        return _wtable(rows) + (_wp("Source: " + b["source"], size=8, color=SLATE) if b.get("source") else "")
    if t == "chart":
        c = norm_chart(b.get("chart", {}))
        if c["kind"] == "forest":
            rows = [["", "Estimate (95% CI)"]] + [[r["label"], f"{r['est']} ({r['lo']}–{r['hi']})"] for r in c.get("rows", [])]
        else:
            rows = [[""] + [s.get("name", "") for s in c["series"]]] + [
                [cat] + [fmt_value(float(s["values"][j]), c) for s in c["series"]]
                for j, cat in enumerate(c["categories"])]
        return _wblock({"type": "table", "rows": rows, "caption": (b.get("caption", "") + " (chart data)").strip(),
                        "source": b.get("source", "")})
    if t == "quote":
        return _wp(b.get("text", ""), italic=True) + _wp("— " + b.get("attribution", ""), size=9, color=SLATE)
    if t == "source":
        return _wp("Source: " + b.get("text", ""), size=8, color=SLATE)
    if t == "image":
        return _wp(f"[Figure: {b.get('caption') or b.get('path', '')} — embed from {b.get('path', '')}]", italic=True, color=SLATE)
    if t == "page_break":
        return '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'
    return _wp(b.get("text", ""))


def report_docx_stdlib(spec: dict, out: Path) -> None:
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" ' \
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    body = [_wp(plain(spec.get("doc_type", "Medical Affairs report")).upper(), size=9, color=TEAL, bold=True),
            _wp(plain(spec.get("title", "")), style="Title")]
    if spec.get("subtitle"):
        body.append(_wp(spec["subtitle"], size=13, color=SLATE))
    meta = " · ".join(str(x) for x in [spec.get("author", ""), spec.get("date", ""), spec.get("doc_id", ""),
                                       ("Version " + str(spec["version"])) if spec.get("version") else ""] if x)
    if meta:
        body.append(_wp(meta, size=9, color=SLATE))
    if spec.get("approval_status"):
        body.append(_wcallout("Approval status", spec["approval_status"], "note"))
    body.append(_wcallout("Document status", DRAFT + ("\n" + SYNTHETIC if spec.get("synthetic") else ""), "safety"))
    if spec.get("summary"):
        body.append(_wcallout("Key takeaways", "\n".join(f"{k}. {plain(x)}" for k, x in enumerate(spec["summary"], 1)), "finding"))
    for k, sec in enumerate(spec.get("sections", []), 1):
        if sec.get("heading"):
            body.append(_wp(f"{k}.  {plain(sec['heading'])}", style="Heading1"))
        for b in _section_blocks(sec):
            body.append(_wblock(b))
    if spec.get("references"):
        body.append(_wp("References", style="Heading1"))
        body += [_wp(f"{k}.  {r}", size=9) for k, r in enumerate(spec["references"], 1)]
    sect = ('<w:sectPr><w:footerReference w:type="default" r:id="rId2"/>'
            '<w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1300" w:right="1350" w:bottom="1250" '
            'w:left="1350" w:header="600" w:footer="500" w:gutter="0"/></w:sectPr>')
    document = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>'
                + "".join(body) + sect + "</w:body></w:document>")
    styles = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {W}>'
              f'<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="{BODY_FONT}" w:hAnsi="{BODY_FONT}" w:cs="{BODY_FONT}"/>'
              f'<w:color w:val="1F2D3A"/><w:sz w:val="21"/></w:rPr></w:rPrDefault>'
              '<w:pPrDefault><w:pPr><w:spacing w:after="120" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
              '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
              f'<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/>'
              f'<w:pPr><w:spacing w:after="160"/></w:pPr><w:rPr><w:rFonts w:ascii="{HEAD_FONT}" w:hAnsi="{HEAD_FONT}"/>'
              f'<w:b/><w:color w:val="{INK}"/><w:sz w:val="52"/></w:rPr></w:style>'
              + "".join(
                  f'<w:style w:type="paragraph" w:styleId="Heading{n}"><w:name w:val="heading {n}"/><w:basedOn w:val="Normal"/>'
                  f'<w:next w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="{b}" w:after="80"/><w:outlineLvl w:val="{n - 1}"/></w:pPr>'
                  f'<w:rPr><w:rFonts w:ascii="{HEAD_FONT}" w:hAnsi="{HEAD_FONT}"/><w:b/><w:color w:val="{c}"/><w:sz w:val="{s}"/></w:rPr></w:style>'
                  for n, s, c, b in ((1, 34, INK, 360), (2, 26, TEAL, 280), (3, 23, INK, 200)))
              + '<w:style w:type="paragraph" w:styleId="ListPara"><w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/>'
                '<w:pPr><w:spacing w:after="60"/><w:ind w:left="360" w:hanging="260"/></w:pPr></w:style>'
                '<w:style w:type="paragraph" w:styleId="ListPara2"><w:name w:val="List Paragraph 2"/><w:basedOn w:val="Normal"/>'
                '<w:pPr><w:spacing w:after="40"/><w:ind w:left="720" w:hanging="260"/></w:pPr></w:style>'
              + "</w:styles>")
    footer = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:ftr {W}><w:p><w:r><w:rPr><w:color w:val="{OXBLOOD}"/>'
              f'<w:sz w:val="14"/></w:rPr><w:t xml:space="preserve">{_x(DRAFT)}   Page </w:t></w:r>'
              '<w:r><w:rPr><w:sz w:val="14"/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r>'
              '<w:r><w:rPr><w:sz w:val="14"/></w:rPr><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
              '<w:r><w:rPr><w:sz w:val="14"/></w:rPr><w:fldChar w:fldCharType="end"/></w:r></w:p></w:ftr>')
    ctypes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
              '<Default Extension="xml" ContentType="application/xml"/>'
              '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
              '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
              '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
              '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>')
    doc_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>'
                '</Relationships>')
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ctypes)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
        z.writestr("word/styles.xml", styles)
        z.writestr("word/footer1.xml", footer)
        z.writestr("word/_rels/document.xml.rels", doc_rels)


# ===========================================================================
# PDF: LibreOffice conversion first (identical to the Word file), reportlab
# direct rendering second.
# ===========================================================================

def office_convert(src: Path, fmt: str = "pdf", outdir: Path | None = None) -> Path | None:
    exe = office_binary()
    if not exe:
        return None
    outdir = outdir or src.parent
    profile = Path(tempfile.gettempdir()) / f"ma_render_lo_{os.getuid() if hasattr(os, 'getuid') else 'u'}"
    env = dict(os.environ)
    env.setdefault("HOME", str(profile))
    try:
        subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless",
                        "--convert-to", fmt, "--outdir", str(outdir), str(src)],
                       capture_output=True, text=True, timeout=240, env=env)
    except Exception:
        return None
    target = outdir / (src.stem + "." + fmt)
    return target if target.exists() and target.stat().st_size > 0 else None


def report_pdf_reportlab(spec: dict, out: Path, workdir: Path) -> list[str]:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak,
                                    Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)
    from reportlab.pdfgen import canvas as rl_canvas

    warnings: list[str] = []
    C = lambda h: colors.HexColor("#" + h)  # noqa: E731
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=14, textColor=C("1F2D3A"), spaceAfter=6)
    small = ParagraphStyle("s", parent=base, fontSize=8, leading=10, textColor=C(SLATE))
    h1 = ParagraphStyle("h1", parent=base, fontName="Times-Bold", fontSize=16, leading=20, textColor=C(INK),
                        spaceBefore=14, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=h1, fontSize=12.5, leading=16, textColor=C(TEAL), spaceBefore=10)
    h3 = ParagraphStyle("h3", parent=h1, fontSize=11, leading=14, spaceBefore=8)
    title = ParagraphStyle("t", parent=base, fontName="Times-Bold", fontSize=26, leading=30, textColor=colors.white)
    cellst = ParagraphStyle("c", parent=base, fontSize=8.5, leading=10.5, spaceAfter=0)
    cellh = ParagraphStyle("ch", parent=cellst, fontName="Helvetica-Bold", textColor=colors.white)

    def P(text, st=base):
        s = []
        for t, b, i in runs(str(text)):
            t = html.escape(t)
            s.append(f"<b>{t}</b>" if b else f"<i>{t}</i>" if i else t)
        return Paragraph("".join(s), st)

    page = A4 if spec.get("page_size", "letter").lower() == "a4" else letter
    avail = page[0] - 1.9 * inch

    class NumberedCanvas(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._saved)
            for st in self._saved:
                self.__dict__.update(st)
                self.setFont("Helvetica", 7.5)
                self.setFillColor(C(OXBLOOD))
                self.drawString(0.95 * inch, 0.5 * inch, DRAFT)
                self.setFillColor(C(SLATE))
                ctrl = " · ".join(str(x) for x in [spec.get("doc_id", ""), spec.get("date", "")] if x)
                self.drawRightString(page[0] - 0.95 * inch, 0.5 * inch,
                                     (ctrl + "   " if ctrl else "") + f"Page {self._pageNumber} of {n}")
                if self._pageNumber > 1:
                    self.drawString(0.95 * inch, page[1] - 0.55 * inch, spec.get("title", "")[:90])
                super().showPage()
            super().save()

    def boxed(flows, fill, accent):
        t = Table([[flows]], colWidths=[avail])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C(fill)),
                               ("LINEBEFORE", (0, 0), (0, -1), 4, C(accent)),
                               ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                               ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        return t

    def table_flow(rows, source="", caption=""):
        n = max(len(r) for r in rows)
        rows = [list(r) + [""] * (n - len(r)) for r in rows]
        lens = [max(len(plain(str(r[c]))) for r in rows) for c in range(n)]
        weights = [min(max(L, 8), 55) for L in lens]
        widths = [avail * w / sum(weights) for w in weights]
        data = [[P(c, cellh if r == 0 else cellst) for c in row] for r, row in enumerate(rows)]
        t = Table(data, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), C(INK)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                 ("LINEBELOW", (0, 0), (-1, -1), 0.4, C("C9D1D8")),
                 ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
        for r in range(2, len(rows), 2):
            style.append(("BACKGROUND", (0, r), (-1, r), C(MIST)))
        t.setStyle(TableStyle(style))
        out_f = []
        if caption:
            out_f.append(P(f"**{caption}**", ParagraphStyle("cap", parent=base, fontSize=9)))
        out_f.append(t)
        if source:
            out_f.append(P("Source: " + source, small))
        return KeepTogether(out_f) if len(rows) < 18 else out_f

    story = []
    cover = [P(plain(spec.get("doc_type", "Medical Affairs report")).upper(),
               ParagraphStyle("k", parent=base, textColor=C("7FC4C2"), fontName="Helvetica-Bold", fontSize=9)),
             P(spec.get("title", ""), title)]
    if spec.get("subtitle"):
        cover.append(P(spec["subtitle"], ParagraphStyle("st", parent=base, textColor=C("C8D4DC"), fontSize=13, leading=17)))
    ct = Table([[cover]], colWidths=[avail])
    ct.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C(INK)), ("LEFTPADDING", (0, 0), (-1, -1), 18),
                            ("TOPPADDING", (0, 0), (-1, -1), 28), ("BOTTOMPADDING", (0, 0), (-1, -1), 28),
                            ("LINEBELOW", (0, 0), (-1, -1), 6, C(TEAL))]))
    story += [ct, Spacer(1, 14)]
    meta = [(k, v) for k, v in [("Prepared by", spec.get("author", "")), ("Date", spec.get("date", "")),
                                ("Version", spec.get("version", "")), ("Document ID", spec.get("doc_id", ""))] if v]
    if meta:
        story.append(table_flow([["Field", "Value"]] + [[k, str(v)] for k, v in meta]))
    if spec.get("approval_status"):
        story.append(boxed([P("**APPROVAL STATUS**", small), P(spec["approval_status"])], MIST, SLATE))
        story.append(Spacer(1, 6))
    story.append(boxed([P("**DOCUMENT STATUS**", small), P(DRAFT + (" " + SYNTHETIC if spec.get("synthetic") else ""))], OX_TINT, OXBLOOD))
    if spec.get("summary"):
        story.append(Spacer(1, 8))
        story.append(boxed([P("**KEY TAKEAWAYS**", small)] + [P(f"**{k}.** {x}") for k, x in enumerate(spec["summary"], 1)], MIST, INK))
    story.append(PageBreak())
    fig = 0
    for k, sec in enumerate(spec.get("sections", []), 1):
        if sec.get("heading"):
            story.append(P(f"{k}.  {sec['heading']}", h1))
        for b in _section_blocks(sec):
            t = b.get("type", "paragraph")
            if t == "paragraph":
                story.append(P(b.get("text", "")))
            elif t == "heading":
                story.append(P(b.get("text", ""), h2 if int(b.get("level", 2)) <= 2 else h3))
            elif t in ("bullets", "numbered"):
                items = []
                for it in b.get("items", []):
                    text = it if isinstance(it, str) else it.get("text", "")
                    items.append(ListItem(P(text), leftIndent=14))
                    if isinstance(it, dict):
                        for sub in it.get("sub", []):
                            items.append(ListItem(P(sub, ParagraphStyle("sb", parent=base, textColor=C(SLATE))), leftIndent=28))
                story.append(ListFlowable(items, bulletType="1" if t == "numbered" else "bullet",
                                          bulletColor=C(TEAL), leftIndent=14))
            elif t == "table":
                f = table_flow(b.get("rows", []), b.get("source", ""), b.get("caption", ""))
                story += f if isinstance(f, list) else [f]
                story.append(Spacer(1, 6))
            elif t == "callout":
                tint, accent, label = CALLOUT_STYLES.get(b.get("style", "note"), CALLOUT_STYLES["note"])
                story.append(boxed([P(f"**{plain(b.get('title') or label).upper()}**", small), P(b.get("text", ""))], tint, accent))
                story.append(Spacer(1, 6))
            elif t == "kpis":
                items = b.get("items", [])[:4]
                cells = [[P(f"**{i.get('value', '')}**", ParagraphStyle('v', parent=base, fontSize=18, leading=22, textColor=C(TEAL))),
                          P(f"**{i.get('label', '')}**", cellst), P(i.get("note", ""), small)] for i in items]
                kt = Table([cells], colWidths=[avail / max(1, len(items))] * len(items))
                kt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C(MIST)), ("LINEABOVE", (0, 0), (-1, 0), 3, C(TEAL)),
                                        ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story.append(kt)
                if b.get("source"):
                    story.append(P("Source: " + b["source"], small))
            elif t == "chart":
                if have("matplotlib"):
                    fig += 1
                    png = workdir / f"pdf_fig_{fig}.png"
                    chart_png(b["chart"], png, width_in=7.2, height_in=3.6)
                    story.append(Image(str(png), width=avail, height=avail / 2))
                    story.append(P(f"**Figure {fig}.** {b.get('caption', '')}  " + (("Source: " + b["source"]) if b.get("source") else ""), small))
                else:
                    warnings.append("matplotlib unavailable — chart rendered as a table in the PDF")
                    c = norm_chart(b["chart"])
                    rows = [[""] + [s.get("name", "") for s in c["series"]]] + [
                        [cat] + [fmt_value(float(s["values"][j]), c) for s in c["series"]]
                        for j, cat in enumerate(c["categories"])]
                    story.append(table_flow(rows, b.get("source", ""), b.get("caption", "")))
            elif t == "image":
                p = Path(b.get("path", ""))
                if p.exists():
                    story.append(Image(str(p), width=avail, height=avail * 0.5, kind="proportional"))
                    if b.get("caption"):
                        story.append(P(b["caption"], small))
            elif t == "quote":
                story.append(P(f"*{b.get('text', '')}*"))
                if b.get("attribution"):
                    story.append(P("— " + b["attribution"], small))
            elif t == "source":
                story.append(P("Source: " + b.get("text", ""), small))
            elif t == "page_break":
                story.append(PageBreak())
    if spec.get("references"):
        story.append(P("References", h1))
        for k, r in enumerate(spec["references"], 1):
            story.append(P(f"{k}. {r}", ParagraphStyle("ref", parent=small, textColor=C("1F2D3A"), fontSize=8.5, leading=11)))
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out), pagesize=page, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                            topMargin=0.85 * inch, bottomMargin=0.85 * inch, title=spec.get("title", ""),
                            author=spec.get("author", "Medical Affairs"), subject=DRAFT)
    doc.build(story, canvasmaker=NumberedCanvas)
    return warnings


# ===========================================================================
# QA and previews
# ===========================================================================

def pdf_to_pngs(pdf: Path, outdir: Path, dpi: int = 60) -> list[Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    pngs: list[Path] = []
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        d = fitz.open(str(pdf))
        for i, page in enumerate(d, 1):
            p = outdir / f"page-{i:02d}.png"
            page.get_pixmap(dpi=dpi).save(str(p))
            pngs.append(p)
        return pngs
    except Exception:
        pass
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", str(dpi), str(pdf), str(outdir / "page")],
                       capture_output=True, timeout=240)
        pngs = sorted(outdir.glob("page*.png"))
    return pngs


def contact_sheet(pngs: list[Path], out: Path, cols: int = 3) -> Path | None:
    try:
        from PIL import Image
    except Exception:
        return None
    if not pngs:
        return None
    ims = [Image.open(p).convert("RGB") for p in pngs]
    w = max(i.width for i in ims)
    h = max(i.height for i in ims)
    rows = math.ceil(len(ims) / cols)
    pad = 12
    sheet = Image.new("RGB", (cols * (w + pad) + pad, rows * (h + pad) + pad), (225, 229, 233))
    for k, im in enumerate(ims):
        r, c = divmod(k, cols)
        sheet.paste(im, (pad + c * (w + pad), pad + r * (h + pad)))
    sheet.save(out)
    return out


def qa_pptx(path: Path) -> list[str]:
    """Structural QA on a finished deck: overflow, off-slide shapes, density,
    tiny fonts, text-only decks, missing sources on data slides."""
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    issues: list[str] = []
    text_only = 0
    content_slides = 0
    for n, slide in enumerate(prs.slides, 1):
        words = 0
        visual = False
        has_source = False
        has_data = False
        for sh in slide.shapes:
            if sh.left is not None and (sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 9525
                                        or sh.top + sh.height > H + 9525):
                issues.append(f"slide {n}: shape '{sh.name}' extends beyond the slide edge")
            if (sh.shape_type == 1 and not (sh.has_text_frame and sh.text_frame.text.strip())
                    and Emu(sh.width).inches * Emu(sh.height).inches >= 1.0
                    and Emu(sh.width).inches < 13.0):
                visual = True  # designed panel: stat card, column, question row
            if sh.shape_type in (13,) or getattr(sh, "has_chart", False) or getattr(sh, "has_table", False):
                visual = True
                if getattr(sh, "has_chart", False) or getattr(sh, "has_table", False):
                    has_data = True
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text
            if sh.name == "Source" or txt.startswith("Source:"):
                has_source = True
                continue
            if sh.name in ("Draft marking", "Page number"):
                continue
            words += len(txt.split())
            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            if not sizes or not txt.strip():
                continue
            pt = max(set(sizes), key=sizes.count)
            if pt < 10:
                issues.append(f"slide {n}: {pt:.0f}pt body text in '{sh.name}' — too small to read")
            paras = [p.text for p in sh.text_frame.paragraphs]
            w_in, h_in = Emu(sh.width).inches, Emu(sh.height).inches
            need = text_height(paras, max(0.5, w_in - 0.2), pt, em=0.5, spacing=1.25, gap_pt=pt * 0.3)
            if need > h_in + 0.25 and h_in > 0.3:
                issues.append(f"slide {n}: text in '{sh.name}' needs ~{need:.1f}in but the box is "
                              f"{h_in:.1f}in — likely overflow")
            if len(slide.shapes) <= 8 and sh.name not in ("Title", "Kicker") and len(paras) >= 3 and pt >= 12:
                pass
        if n > 1:
            content_slides += 1
            if not visual and words > 25:
                text_only += 1
            if words > 120:
                issues.append(f"slide {n}: {words} words — a slide is not a document; split it or "
                              f"move detail to notes")
            if has_data and not has_source:
                issues.append(f"slide {n}: chart/table without a 'Source:' line")
    if content_slides and text_only / content_slides > 0.6:
        issues.append(f"{text_only}/{content_slides} content slides are text only — convert numbers "
                      f"into 'stats' or 'chart' slides and comparisons into 'two_column'")
    return issues


def preview(path: Path, outdir: Path) -> Path | None:
    """Render a deck or document to PNG pages + one contact sheet for visual review."""
    outdir.mkdir(parents=True, exist_ok=True)
    pdf = path if path.suffix.lower() == ".pdf" else office_convert(path, "pdf", outdir)
    if not pdf:
        return None
    pngs = pdf_to_pngs(pdf, outdir)
    return contact_sheet(pngs, outdir / "contact-sheet.png")


# ===========================================================================
# Public entry points
# ===========================================================================

def _notice(kind: str, delivered: list[Path], missing: str, how: str) -> None:
    print(f"\n⚠ DEGRADED OUTPUT — {missing}\n")
    print("   Delivered:      " + "\n                   ".join(str(p) for p in delivered))
    print(f"   To get the full {kind}: {how}")
    print("   Every citation, design statement, denominator, safety line and the draft\n"
          "   marking are in the delivered file; only layout fidelity changed.\n")


def render_deck(spec: dict, out: Path, *, pdf: bool = True, install: bool | None = None,
                strict_approval: bool = False, preview_dir: Path | None = None) -> dict:
    errors, warnings = validate_deck(spec, strict_approval=strict_approval)
    result = {"errors": errors, "warnings": warnings, "files": []}
    if errors:
        return result
    status = ensure(["pptx", "matplotlib"], install=install)
    out = Path(out)
    workdir = out.parent / f".{out.stem}_assets"
    if status["pptx"]:
        workdir.mkdir(parents=True, exist_ok=True)
        builder = DeckBuilder(spec, workdir)
        result["warnings"] += builder.build(out)
        result["files"].append(out)
        try:
            result["warnings"] += [w for w in qa_pptx(out) if w not in result["warnings"]]
        except Exception as exc:
            result["warnings"].append(f"QA skipped: {exc}")
        if pdf:
            p = office_convert(out, "pdf")
            if p:
                result["files"].append(p)
            else:
                result["notes"] = ["PDF copy needs LibreOffice (soffice); the .pptx is complete. "
                                   "Export to PDF from PowerPoint/Keynote/Google Slides if needed."]
        if preview_dir:
            sheet = preview(out, Path(preview_dir))
            if sheet:
                result["preview"] = sheet
        spec_copy = out.with_suffix(".json")
        spec_copy.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
        result["spec"] = spec_copy
        return result
    html_path = out.with_suffix(".html")
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(deck_html(spec), encoding="utf-8")
    spec_path = out.with_suffix(".json")
    spec_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    result["files"] += [html_path, spec_path]
    result["degraded"] = True
    _notice("deck", [html_path, spec_path],
            "python-pptx is not available and could not be installed.",
            f"pip install python-pptx  then  python3 {Path(__file__).name} deck {spec_path} --out {out}. "
            f"The HTML prints one 16:9 slide per page (Ctrl/Cmd-P → Save as PDF).")
    return result


def render_report(spec: dict, out: Path, *, pdf: bool = True, install: bool | None = None,
                  preview_dir: Path | None = None) -> dict:
    errors, warnings = validate_report(spec)
    result = {"errors": errors, "warnings": warnings, "files": []}
    if errors:
        return result
    out = Path(out)
    if out.suffix.lower() == ".pdf":
        pdf_only, docx_out = True, out.with_suffix(".docx")
    else:
        pdf_only, docx_out = False, out
    status = ensure(["docx", "matplotlib"] + (["reportlab"] if pdf else []), install=install)
    workdir = out.parent / f".{out.stem}_assets"
    workdir.mkdir(parents=True, exist_ok=True)
    degraded = False
    if status["docx"]:
        result["warnings"] += ReportBuilder(spec, workdir).build(docx_out)
    else:
        report_docx_stdlib(spec, docx_out)
        degraded = True
    result["files"].append(docx_out)
    if pdf or pdf_only:
        p = office_convert(docx_out, "pdf")
        if not p and status.get("reportlab"):
            p = out.with_suffix(".pdf")
            result["warnings"] += report_pdf_reportlab(spec, p, workdir)
        if p:
            result["files"].append(p)
        else:
            h = out.with_suffix(".html")
            h.write_text(report_html(spec), encoding="utf-8")
            result["files"].append(h)
            result.setdefault("notes", []).append(
                f"PDF needs reportlab or LibreOffice; wrote print-ready {h.name} (Ctrl/Cmd-P → Save as PDF).")
    if preview_dir:
        target = next((f for f in result["files"] if str(f).endswith(".pdf")), None)
        if target:
            sheet = preview(target, Path(preview_dir))
            if sheet:
                result["preview"] = sheet
    spec_path = out.with_suffix(".json")
    spec_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    result["spec"] = spec_path
    if degraded:
        result["degraded"] = True
        _notice("Word document", result["files"],
                "python-docx is not available and could not be installed; a plainer .docx "
                "was written with the standard library.",
                f"pip install python-docx matplotlib  then  python3 {Path(__file__).name} report {spec_path} --out {docx_out}")
    return result


# ===========================================================================
# Worked examples (fictional product; synthetic numbers)
# ===========================================================================

EXAMPLE_DECK = {
    "title": "Sustaining benefit in routine practice: what the evidence does and does not show",
    "subtitle": "Advisory board pre-read — evidence gaps and education priorities",
    "kicker": "Advisory board · pre-read",
    "presenter": "Medical Affairs",
    "date": "2026-09-12",
    "jurisdiction": "US",
    "synthetic": True,
    "approval_status": "NORVANTIB is approved in the US for adults with relapsed/refractory "
                       "multiple myeloma after ≥4 prior lines. Other uses discussed are investigational.",
    "slides": [
        {"type": "agenda", "title": "What we need from you today",
         "items": ["Where the evidence is strong — and where it is not",
                   "The persistence gap in routine practice",
                   "Three evidence bets we are weighing",
                   "Your challenge: what would change your practice?"]},
        {"type": "section", "title": "The evidence base", "subtitle": "Efficacy is established on treatment; durability is not"},
        {"type": "stats", "title": "On treatment, responses are frequent and durable to 18 months",
         "design": "Single-arm phase 1/2 · n=165 · triple-class exposed",
         "items": [{"value": "63%", "label": "Overall response rate", "note": "95% CI 55.2–70.4"},
                   {"value": "18.4 mo", "label": "Median duration of response", "note": "95% CI 14.9–NE; immature"},
                   {"value": "0.6%", "label": "Grade ≥3 CRS", "note": "Any grade 72.1%", "style": "safety"}],
         "citation": "Example A, et al. J Example Med. 2024;12(3):100-110. PMID 00000000",
         "takeaway": "Single-arm data: no comparative inference against current standard of care."},
        {"type": "chart", "title": "Response is lower after prior BCMA-directed therapy",
         "design": "Single-arm phase 1/2 · subgroup analysis · n=165",
         "chart": {"kind": "column", "categories": ["BCMA-naïve (n=112)", "Prior ADC (n=31)", "Prior CAR-T (n=22)"],
                   "values": [71.4, 48.4, 40.9], "unit": "%", "highlight": 0, "y_max": 100,
                   "y_label": "Overall response rate"},
         "bullets": ["Subgroups were not powered for comparison", "Prior CAR-T group is small (n=22)",
                     "**Question for the board:** does this match your experience?"],
         "citation": "Example A, et al. J Example Med. 2024;12(3):100-110. PMID 00000000",
         "tier": "peer-reviewed"},
        {"type": "table", "title": "Safety profile: CRS is common but rarely severe",
         "design": "Single-arm phase 1/2 · n=165 · median follow-up 14.2 months",
         "rows": [["Adverse event", "Any grade", "Grade ≥3", "Management"],
                  ["Cytokine release syndrome", "72.1%", "0.6%", "Step-up dosing; tocilizumab"],
                  ["Infections", "64.2%", "38.8%", "Prophylaxis; IVIG per guidance"],
                  ["Neutropenia", "51.5%", "47.9%", "G-CSF as needed"],
                  ["ICANS", "3.0%", "0%", "Monitoring during step-up"]],
         "highlight_row": 2,
         "citation": "Example A, et al. J Example Med. 2024;12(3):100-110. PMID 00000000"},
        {"type": "two_column", "title": "What is known — and what is not",
         "left": {"heading": "Established", "bullets": ["ORR 63% in triple-class exposed (single-arm)",
                                                        "CRS predominantly grade 1–2 with step-up dosing"]},
         "right": {"heading": "Unknown", "bullets": ["Comparative effectiveness vs current standard",
                                                     "Outcomes beyond 24 months", "Sequencing after CAR-T"]}},
        {"type": "questions", "title": "Questions for the group",
         "questions": ["What would you need to see before using this earlier in the pathway?",
                       "Where does your practice diverge from the label, and why?",
                       "Which evidence gap, if closed, would most change your decisions?"]},
        {"type": "references", "title": "References",
         "references": ["Example A, et al. J Example Med. 2024;12(3):100-110. PMID 00000000"]},
    ],
}

EXAMPLE_REPORT = {
    "title": "Field insights: what leadership should know this quarter",
    "subtitle": "Oncology · relapsed/refractory multiple myeloma",
    "doc_type": "Leadership brief",
    "author": "Medical Affairs — Field Insights",
    "date": "2026-09-12", "version": "0.1", "doc_id": "FI-2026-Q3-001",
    "synthetic": True,
    "approval_status": "NORVANTIB is approved in the US for adults with relapsed/refractory multiple "
                       "myeloma after ≥4 prior lines. Other uses are investigational.",
    "summary": ["Infection management, not CRS, is now the dominant HCP concern (18 of 40 insights).",
                "Community sites report delays of 3–5 weeks for step-up dosing capacity.",
                "One simulated safety finding was routed to the workshop intake exercise."],
    "sections": [
        {"heading": "Bottom line",
         "blocks": [{"type": "callout", "style": "finding", "title": "Key finding",
                     "text": "HCP concern has shifted from CRS to infections; our education plan still leads with CRS."},
                    {"type": "kpis", "items": [{"value": "40", "label": "Insights screened", "note": "12 MSLs, 6 weeks"},
                                               {"value": "18", "label": "Infection-related", "note": "45% of total"},
                                               {"value": "1", "label": "Simulated safety finding", "note": "Routed to intake", "style": "safety"}],
                     "source": "SYN:field-observations.csv, rows 1–40"}]},
        {"heading": "Insight detail",
         "blocks": [{"type": "chart", "caption": "Insight themes (n=40)",
                     "chart": {"kind": "bar", "categories": ["Infections", "Step-up capacity", "Sequencing", "CRS", "Other"],
                               "values": [18, 9, 6, 4, 3], "decimals": 0, "highlight": 0},
                     "source": "SYN:field-observations.csv"},
                    {"type": "table", "caption": "Ranked insights and actions",
                     "rows": [["ID", "Insight", "n", "Confidence", "Owner"],
                              ["INS-001", "Infection prophylaxis practice varies widely", "18", "High", "Medical Education Lead"],
                              ["INS-002", "Step-up dosing delayed by bed capacity", "9", "Moderate", "Field Medical Director"],
                              ["INS-003", "Sequencing after CAR-T is unclear to HCPs", "6", "Low — weak signal", "Evidence Generation Lead"]],
                     "source": "SYN:field-observations.csv"},
                    {"type": "callout", "style": "limitation", "title": "Limitation",
                     "text": "Insights reflect MSL-selected conversations and are not a representative sample of prescribers."}]},
        {"heading": "Recommended actions",
         "blocks": [{"type": "numbered", "items": ["Re-sequence the education plan to lead with infection management (owner: Medical Education Lead, by Q4).",
                                                   "Map step-up capacity at the top 20 community sites (owner: Field Medical Director).",
                                                   "Scope a sequencing evidence review before any new study is proposed."]}]},
    ],
    "references": ["Example A, et al. J Example Med. 2024;12(3):100-110. PMID 00000000"],
}


# ===========================================================================
# CLI
# ===========================================================================

def _print_result(res: dict, kind: str) -> int:
    for e in res.get("errors", []):
        print(f"  BLOCKED  {e}")
    if res.get("errors"):
        print(f"\n{len(res['errors'])} blocking problem(s) — fix the spec and re-run. These are "
              f"compliance and completeness rules, not style preferences.")
        return 1
    for f in res.get("files", []):
        print(f"  wrote    {f}")
    if res.get("preview"):
        print(f"  preview  {res['preview']}   ← open this image and look at every page before delivering")
    for w in res.get("warnings", []):
        print(f"  check    {w}")
    for n in res.get("notes", []):
        print(f"  note     {n}")
    if not res.get("degraded"):
        print(f"\n{kind} complete. The draft marking is on every page; the reviewer removes it, not you.")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="ma_render", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd")
    for name in ("deck", "report", "pdf"):
        p = sub.add_parser(name)
        p.add_argument("spec", help="JSON spec or markdown draft (.md)")
        p.add_argument("--out", required=True)
        p.add_argument("--no-pdf", action="store_true", help="skip the PDF companion")
        p.add_argument("--no-install", action="store_true")
        p.add_argument("--preview", help="folder for PNG page renders + contact sheet")
        p.add_argument("--strict-approval", action="store_true")
    q = sub.add_parser("qa"); q.add_argument("file"); q.add_argument("--preview")
    e = sub.add_parser("example"); e.add_argument("kind", choices=["deck", "report"])
    m = sub.add_parser("md2spec"); m.add_argument("file"); m.add_argument("--kind", choices=["deck", "report"], default="report")
    c = sub.add_parser("chart"); c.add_argument("spec"); c.add_argument("--out", required=True)
    b = sub.add_parser("bootstrap"); b.add_argument("--no-install", action="store_true")
    sub.add_parser("version")
    args = ap.parse_args(argv)

    if args.cmd == "version":
        print(VERSION); return 0
    if args.cmd == "example":
        print(json.dumps(EXAMPLE_DECK if args.kind == "deck" else EXAMPLE_REPORT, indent=2, ensure_ascii=False))
        return 0
    if args.cmd == "bootstrap":
        caps = capabilities(install=False if args.no_install else None)
        print(json.dumps(caps, indent=2, default=str))
        missing = [k for k, v in caps["packages"].items() if not v]
        if missing:
            print(f"\nStill missing: {', '.join(missing)}. Run: {sys.executable} -m pip install "
                  + " ".join(PACKAGES[m] for m in missing))
        if not caps["office_converter"]:
            print("LibreOffice not found: .pptx/.docx are complete; PDFs come from reportlab. "
                  "Install LibreOffice for PDF copies of decks and for page previews.")
        return 0
    if args.cmd == "md2spec":
        print(json.dumps(load_spec(Path(args.file), args.kind), indent=2, ensure_ascii=False))
        return 0
    if args.cmd == "chart":
        ensure(["matplotlib"])
        chart_png(json.loads(Path(args.spec).read_text(encoding="utf-8")), Path(args.out))
        print(f"  wrote    {args.out}")
        return 0
    if args.cmd == "qa":
        ensure(["pptx"])
        path = Path(args.file)
        issues = qa_pptx(path) if path.suffix.lower() == ".pptx" else []
        for i in issues:
            print(f"  check    {i}")
        if args.preview:
            sheet = preview(path, Path(args.preview))
            print(f"  preview  {sheet}" if sheet else "  preview  unavailable (needs LibreOffice + pymupdf or pdftoppm)")
        print(f"\n{len(issues)} issue(s).")
        return 0
    if args.cmd in ("deck", "report", "pdf"):
        src = Path(args.spec)
        if not src.exists():
            print(f"No such file: {src}", file=sys.stderr); return 2
        kind = "deck" if args.cmd == "deck" else "report"
        try:
            spec = load_spec(src, kind)
        except json.JSONDecodeError as exc:
            print(f"Spec is not valid JSON: {exc}", file=sys.stderr); return 2
        install = False if args.no_install else None
        prev = Path(args.preview) if args.preview else None
        if args.cmd == "deck":
            res = render_deck(spec, Path(args.out), pdf=not args.no_pdf, install=install,
                              strict_approval=args.strict_approval, preview_dir=prev)
            return _print_result(res, "Deck")
        out = Path(args.out)
        if args.cmd == "pdf" and out.suffix.lower() != ".pdf":
            out = out.with_suffix(".pdf")
        res = render_report(spec, out, pdf=not args.no_pdf or args.cmd == "pdf", install=install, preview_dir=prev)
        return _print_result(res, "Report")
    ap.print_help()
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
